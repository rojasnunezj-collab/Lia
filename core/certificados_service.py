# ====================================================================
# --- MÓDULO CORE: SISTEMA DE CERTIFICADOS IA (LÍA) ---
# ====================================================================
import os
import io
import re
import json
import time
import asyncio
import unicodedata
from datetime import datetime, timezone, timedelta
from typing import Dict, List, Any, Optional, Tuple

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.oxml.simpletypes import ST_TwipsMeasure, Twips
from docxtpl import DocxTemplate
from pypdf import PdfReader, PdfWriter
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseUpload, MediaIoBaseDownload
from google import genai
from google.genai import types

from config.settings import logger, SHEET_ID, DRIVE_FOLDER_ID, PROJECT_ID, REGION_ESTABLE, MODEL_NAME, FALLBACK_MODELS
from core.sheets_client import obtener_credenciales
from utils.helpers import liberar_memoria_proceso

# Zona horaria Perú (UTC-5)
PET = timezone(timedelta(hours=-5))

# ====================================================================
# --- BLOQUE 1: CONSTANTES Y CONFIGURACIÓN DRIVE / SHEETS ---
# ====================================================================
ID_SHEET_CONTROL = SHEET_ID or "14As5bCpZi56V5Nq1DRs0xl6R1LuOXLvRRoV26nI50NU"
ID_SHEET_REPOSITORIO = SHEET_ID or "14As5bCpZi56V5Nq1DRs0xl6R1LuOXLvRRoV26nI50NU"

CARPETA_PLANTILLAS_NORMALES = '1EwbYAbyv2uMsSn0yXZd0vTPuCoPMbzKs'
CARPETA_PLANTILLAS_MODELO = '1_kY1h6PwlhDPl8BjG7u0fbGMT1AWyn3n'

CARPETAS_DESTINO: Dict[str, Dict[str, str]] = {
    "EPMI S.A.C.": {
        "Comercialización": "1NZc-nfGHw5bnkCAv0TdQYW_bPM_UkKC-",
        "Disposición Final 1": "12PMJ1d-CSWo64m7aNQRQj2yGHFdp9B9S",
        "Disposición Final 2": "12PMJ1d-CSWo64m7aNQRQj2yGHFdp9B9S"
    },
    "INECOVE S.A.C.": {
        "Comercialización": "1NZc-nfGHw5bnkCAv0TdQYW_bPM_UkKC-",
        "Disposición Final 1": "12PMJ1d-CSWo64m7aNQRQj2yGHFdp9B9S",
        "Disposición Final 2": "12PMJ1d-CSWo64m7aNQRQj2yGHFdp9B9S"
    }
}

PLANTILLAS_FALLBACK: Dict[str, Dict[str, str]] = {
    "EPMI S.A.C.": {
        "Comercialización": "1d09vmlBlW_4yjrrz5M1XM8WpCvzTI4f11pERDbxFvNE",
        "Disposición Final 1": "1QqqVJ2vCiAjiKKGt_zEpaImUB-q3aRurSiXjMEU--eg",
        "Disposición Final 2": "1fpdZef3Fe3tl00yAuM0Cehx2_o3AusrErcOJisBtdBM"
    },
    "INECOVE S.A.C.": {
        "Comercialización": os.getenv("TEMPLATE_INECOVE_ID", "1JbTdRtinFfTnjDLKmrm7Dvd6DI3p0_kqn_fwwOoIBgU"),
        "Disposición Final 1": os.getenv("TEMPLATE_INECOVE_PELIGROSO_ID", "1W-HyVSivqug13gBRBclBuICAOSBUHm1WN5cnqtMQcZY"),
        "Disposición Final 2": os.getenv("TEMPLATE_INECOVE_PELIGROSO_ID", "1W-HyVSivqug13gBRBclBuICAOSBUHm1WN5cnqtMQcZY")
    }
}

MESES_ES = {
    "01": "Enero", "02": "Febrero", "03": "Marzo", "04": "Abril",
    "05": "Mayo", "06": "Junio", "07": "Julio", "08": "Agosto",
    "09": "Septiembre", "10": "Octubre", "11": "Noviembre", "12": "Diciembre"
}

# ====================================================================
# --- BLOQUE 2: PARCHE TÉCNICO DOCX (TWIPS MEASURE) ---
# ====================================================================
try:
    _orig_convert_from_xml = ST_TwipsMeasure.convert_from_xml

    @classmethod
    def _patch_convert_from_xml(cls, str_value):
        try:
            return Twips(int(str_value))
        except ValueError:
            try:
                return Twips(int(float(str_value)))
            except Exception:
                return _orig_convert_from_xml(str_value)

    ST_TwipsMeasure.convert_from_xml = _patch_convert_from_xml
except Exception as _e:
    logger.debug(f"Aviso parche TwipsMeasure: {_e}")

# ====================================================================
# --- BLOQUE 3: CONEXIÓN A SERVICIOS GOOGLE WORKSPACE ---
# ====================================================================
_cached_drive_service = None
_cached_sheets_service = None

def obtener_servicios_google(forzar_refresco: bool = False):
    """
    Retorna clientes autenticados de Drive y Sheets usando un patrón Singleton reutilizable.
    Evita instanciar decenas de veces los metadatos de discovery de Google en memoria.
    """
    global _cached_drive_service, _cached_sheets_service
    if not forzar_refresco and _cached_drive_service is not None and _cached_sheets_service is not None:
        return _cached_drive_service, _cached_sheets_service

    creds = obtener_credenciales()
    if not creds:
        raise ValueError("No se pudieron obtener credenciales válidas de Google.")
    _cached_drive_service = build('drive', 'v3', credentials=creds, cache_discovery=False)
    _cached_sheets_service = build('sheets', 'v4', credentials=creds, cache_discovery=False)
    return _cached_drive_service, _cached_sheets_service

# ====================================================================
# --- BLOQUE 4: UTILIDADES DE FORMATEO Y CADENAS ---
# ====================================================================
def limpiar_monto(valor: Any) -> float:
    """Convierte string numérico a float limpiando comas y espacios."""
    if not valor:
        return 0.0
    s = str(valor).strip().replace(',', '.')
    if s.count('.') > 1:
        parts = s.split('.')
        s = "".join(parts[:-1]) + '.' + parts[-1]
    s = re.sub(r'[^\d.]', '', s)
    try:
        return float(s)
    except Exception:
        return 0.0

def formato_inteligente(valor: Any) -> str:
    """Formatea números quitando decimales superfluos (100.0 -> 100)."""
    try:
        f = float(valor)
        return str(int(f)) if f.is_integer() else f"{f}"
    except Exception:
        return str(valor)

def normalizar_fecha(fecha_str: str) -> str:
    """Normaliza formatos de fecha a dd/mm/yyyy."""
    if not fecha_str:
        return datetime.now(PET).strftime("%d/%m/%Y")
    for fmt in ["%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"]:
        try:
            return datetime.strptime(fecha_str.strip(), fmt).strftime("%d/%m/%Y")
        except Exception:
            continue
    return str(fecha_str).strip()

def normalizar_texto_sin_tildes(texto: Any) -> str:
    """Elimina tildes y diacríticos para comparaciones robustas."""
    t = str(texto or '').lower()
    return ''.join(c for c in unicodedata.normalize('NFD', t) if unicodedata.category(c) != 'Mn')

def limpiar_descripcion(texto: Any) -> str:
    """Limpia prefijos de residuos como 'VEN - AMB -'."""
    if not texto:
        return ""
    return re.sub(r'VEN\s*-\s*AMB\s*-\s*', '', str(texto).strip(), flags=re.IGNORECASE).strip().upper()

def formatear_titulo_oficial(titulo_raw: str, tipo_cod: str = "COM") -> str:
    """
    Formatea el título del certificado con el salto de línea oficial (\n)
    para que docxtpl inyecte <w:br/> y se renderice en Word/PDF en dos líneas
    perfectamente balanceadas y centradas con la tipografía corporativa.
    """
    t = str(titulo_raw or "").strip()
    if not t or t.upper() in ["S/D", "NONE", "CERTIFICADO DE OPERACIÓN", "CERTIFICADO DE OPERACION"]:
        if tipo_cod == "COM":
            return "CERTIFICADO DE RECOLECCIÓN, TRANSPORTE Y \nCOMERCIALIZACIÓN DE RESIDUOS SÓLIDOS"
        else:
            return "CERTIFICADO DE RECOLECCIÓN Y TRANSPORTE \nDE RESIDUOS SÓLIDOS PELIGROSOS"

    if "\n" in t:
        return t.replace("\r", "").strip()

    t_norm = normalizar_texto_sin_tildes(t).upper()

    if "COMERCIALIZACION" in t_norm:
        m = re.search(r'(CERTIFICADO DE RECOLECCI[OÓ]N,\s*TRANSPORTE\s+Y)\s+(COMERCIALIZACI[OÓ]N.*)', t, re.IGNORECASE)
        if m:
            return f"{m.group(1).strip()} \n{m.group(2).strip()}"
        return "CERTIFICADO DE RECOLECCIÓN, TRANSPORTE Y \nCOMERCIALIZACIÓN DE RESIDUOS SÓLIDOS"

    if "PELIGROSO" in t_norm and "NO PELIGROSO" not in t_norm:
        m = re.search(r'(CERTIFICADO DE RECOLECCI[OÓ]N\s+Y\s+TRANSPORTE)\s+(DE\s+RESIDUOS.*)', t, re.IGNORECASE)
        if m:
            return f"{m.group(1).strip()} \n{m.group(2).strip()}"
        return "CERTIFICADO DE RECOLECCIÓN Y TRANSPORTE \nDE RESIDUOS SÓLIDOS PELIGROSOS"

    if "NO PELIGROSO" in t_norm:
        m = re.search(r'(CERTIFICADO DE RECOLECCI[OÓ]N\s+Y\s+TRANSPORTE)\s+(DE\s+RESIDUOS.*)', t, re.IGNORECASE)
        if m:
            return f"{m.group(1).strip()} \n{m.group(2).strip()}"
        return "CERTIFICADO DE RECOLECCIÓN Y TRANSPORTE \nDE RESIDUOS SÓLIDOS NO PELIGROSOS"

    if "TRANSPORTE Y" in t.upper():
        idx = t.upper().find("TRANSPORTE Y") + len("TRANSPORTE Y")
        return f"{t[:idx].strip()} \n{t[idx:].strip()}"
    elif "TRANSPORTE" in t.upper() and "DE RESIDUOS" in t.upper():
        idx = t.upper().find("TRANSPORTE") + len("TRANSPORTE")
        return f"{t[:idx].strip()} \n{t[idx:].strip()}"

    return t

def limpiar_numero_guia(val: Any) -> str:
    """
    Limpia y estandariza el número de guía para la Mini App y Certificados,
    eliminando ceros a la izquierda del correlativo y de la serie si aplica.
    Ejemplos:
      'tr16-00012'        -> 'TR16-12'
      'TR16 - N° 0004452' -> 'TR16-4452'
      'TR01-0000026'      -> 'TR01-26'
      'EG07 - 00000074'   -> 'EG07-74'
      'T002-0000707'      -> 'T002-707'
      'TO01 - 0000893'    -> 'TO01-893'
      '016-5809'          -> '16-5809'
      '001-2607-0005'     -> '1-2607-5'
      '0005530'           -> '5530'
      'TR1600012'         -> 'TR16-12'
    """
    if not val:
        return ""
    s = str(val).strip().upper()
    if s in ["", "NONE", "NAN", "S/D", "S/G", "SIN DATOS"]:
        return ""

    # Remover prefijos como GUIA DE REMISION, N°, Nº, etc.
    s = re.sub(r'^(?:GU[IÍ]A(?:\s+DE\s+REMISI[OÓ]N)?\s*(?:ELECTR[OÓ]NICA)?\s*:?\s*)', '', s)
    s = re.sub(r'(?:N[°º\.]|NUMERO|NRO\.?)\s*', '', s)

    # Si no tiene dígitos (ej. 'SIN GUIA', 'NO TIENE', 'BALANZA'), conservar texto
    if not any(c.isdigit() for c in s):
        return s

    # Si son exclusivamente dígitos
    if s.isdigit():
        return s.lstrip('0') or '0'

    # Normalizar separadores (- o /) y espacios entre serie y correlativo
    s_norm = re.sub(r'\s*[-/]\s*', '-', s)
    s_norm = re.sub(r'([A-Z0-9]+)\s+([0-9]+)', r'\1-\2', s_norm)

    if '-' in s_norm:
        partes = s_norm.split('-')
        partes_norm = [p.lstrip('0') or '0' for p in partes]
        return "-".join(partes_norm)

    # Caso: Serie alfanumérica de 4 caracteres con letras iniciales y correlativo pegado (ej. TR1600012 o T00100045)
    m2 = re.match(r'^([A-Z][A-Z0-9]{3})([0-9]+)$', s)
    if m2:
        serie = m2.group(1).strip()
        num = m2.group(2).lstrip('0') or '0'
        return f"{serie}-{num}"

    return s

def es_formato_guia(valor: Any) -> bool:
    """Determina si un valor tiene formato característico de Guía de Remisión."""
    if not valor:
        return False
    s = str(valor).strip().upper()
    if s in ['', 'NONE', 'NAN', 'S/N', 'SIN DATOS']:
        return False
    s = re.sub(r'^(?:GU[IÍ]A(?:\s+DE\s+REMISI[OÓ]N)?\s*(?:ELECTR[OÓ]NICA)?\s*:?\s*)', '', s)
    s = re.sub(r'(?:N[°º\.]|NUMERO|NRO\.?)\s*', '', s)
    if '-' in s:
        partes = s.split('-', 1)
        serie = partes[0].strip()
        correlativo = partes[1].strip()
        if re.match(r'^(T\d{2,3}|EG\d{2}|E\d{3}|V\d{3}|GR\d{2}|TR\d{2}|TO\d{2}|[A-Z0-9]{2,6})$', serie) and correlativo.isdigit():
            return True
    return False

def es_formato_placa(valor: Any) -> bool:
    """Determina si un valor corresponde al patrón de una placa vehicular peruana."""
    if not valor:
        return False
    s = str(valor).strip().upper()
    if s in ['', 'NONE', 'NAN', 'S/N', 'SIN DATOS']:
        return False
    s_limpio = s.replace('-', '').replace(' ', '')
    if len(s_limpio) not in [5, 6]:
        return False
    if re.match(r'^(T\d{3}|EG\d{2}|E\d{3}|V\d{3}|GR\d{2}|TR\d{2}|TO\d{2})$', s_limpio[:4]):
        return False
    if re.match(r'^[A-Z0-9]{3}[A-Z0-9]{3}$', s_limpio):
        tiene_letras = bool(re.search(r'[A-Z]', s_limpio))
        tiene_numeros = bool(re.search(r'\d', s_limpio))
        if tiene_letras and tiene_numeros:
            return True
    if re.match(r'^[A-Z]{2}\d{4}$', s_limpio):
        return True
    return False

def extraer_id_drive(valor: Any) -> Optional[str]:
    """Extrae de forma robusta el fileId de URLs de Google Drive o Docs."""
    if not valor:
        return None
    val_str = str(valor).strip()
    # Si parece nombre de archivo o prefijo de certificado/guía, no es ID
    if val_str.upper().startswith(('CERT-', 'TR', 'EG', 'GR', 'T0', 'F0', 'EPMI')) or any(val_str.lower().endswith(ext) for ext in ['.pdf', '.docx', '.xlsx', '.doc']):
        return None

    patrones = [
        r'drive\.google\.com/file/d/([a-zA-Z0-9_-]+)',
        r'docs\.google\.com/document/d/([a-zA-Z0-9_-]+)',
        r'drive\.google\.com/open\?id=([a-zA-Z0-9_-]+)',
        r'id=([a-zA-Z0-9_-]+)'
    ]
    for p in patrones:
        m = re.search(p, val_str)
        if m:
            return m.group(1)

    if re.match(r'^[a-zA-Z0-9_-]{25,65}$', val_str) and not re.match(r'^[A-Z0-9_]+-[A-Z0-9_]+', val_str):
        return val_str

    return None

def parsear_celda_enlace_o_nombre(valor: Any) -> Tuple[str, str]:
    """
    Retorna (url, nombre_archivo) a partir de:
    - =HYPERLINK("url"; "nombre") o =HYPERLINK("url", "nombre")
    - URL cruda (https://...)
    - Nombre de archivo plano (CERT-COM-...)
    """
    if not valor:
        return "", ""
    s = str(valor).strip()
    if s.startswith('#'):
        return "", ""
    m = re.search(r'=HYPERLINK\(\s*"([^"]+)"\s*(?:[;,]\s*"([^"]+)")?\s*\)', s, re.IGNORECASE)
    if m:
        url = m.group(1).strip()
        nombre = m.group(2).strip() if m.group(2) else ""
        return url, nombre
    if s.startswith(('http://', 'https://')):
        return s, ""
    return "", s

def obtener_link_archivo_drive(servicio_drive, nombre_o_id: str) -> Optional[str]:
    """Obtiene webViewLink público/corporativo para un archivo por ID o nombre."""
    if not servicio_drive or not nombre_o_id:
        return None
    s_nom = str(nombre_o_id).strip()
    if s_nom.startswith('#') or s_nom.upper() in ['#ERROR!', '#REF!', '#VALUE!', '#N/A']:
        return None
    try:
        f_id = extraer_id_drive(s_nom)
        if f_id:
            try:
                meta = servicio_drive.files().get(
                    fileId=f_id, fields='webViewLink', supportsAllDrives=True
                ).execute()
                if meta.get('webViewLink'):
                    return meta.get('webViewLink')
            except Exception:
                pass

        query = f"name contains '{s_nom}' and trashed = false"
        res = servicio_drive.files().list(
            q=query, spaces='drive', fields='files(id, webViewLink)', supportsAllDrives=True,
            orderBy='modifiedTime desc'
        ).execute()
        files = res.get('files', [])
        return files[0].get('webViewLink') if files else None
    except Exception as e:
        logger.error(f"Error resolviendo link en Drive para '{s_nom}': {e}")
        return None

def descargar_archivo_drive_por_id_o_nombre(nombre_o_id: str, drive_service=None) -> Optional[bytes]:
    """
    Descarga el contenido binario de un archivo de Google Drive
    identificado por ID, URL o nombre de archivo asegurando la última revisión.
    """
    if not nombre_o_id:
        return None
    s_id = str(nombre_o_id).strip()
    if s_id.startswith('#') or s_id.upper() in ['#ERROR!', '#REF!', '#VALUE!', '#N/A']:
        return None
    if not drive_service:
        drive_service, _ = obtener_servicios_google()
    if not nombre_o_id:
        return None

    archivo_id = extraer_id_drive(nombre_o_id)
    es_posible_id = bool(archivo_id) or bool(re.match(r'^[a-zA-Z0-9_-]{25,65}$', str(nombre_o_id).strip()))
    file_id_final = archivo_id or (str(nombre_o_id).strip() if es_posible_id else None)

    try:
        if file_id_final:
            meta = drive_service.files().get(
                fileId=file_id_final, fields='id, name, mimeType', supportsAllDrives=True
            ).execute()
            mime_type = meta.get('mimeType', '')
            if mime_type == 'application/vnd.google-apps.document':
                req = drive_service.files().export_media(fileId=file_id_final, mimeType='application/pdf')
            else:
                try:
                    rev_res = drive_service.revisions().list(
                        fileId=file_id_final, fields='revisions(id, modifiedTime)'
                    ).execute()
                    revs = rev_res.get('revisions', [])
                    if revs:
                        latest_rev = sorted(revs, key=lambda r: r.get('modifiedTime', ''))[-1]
                        logger.info(f"Descargando última revisión {latest_rev.get('id')} ({latest_rev.get('modifiedTime')}) para {file_id_final}")
                        req = drive_service.revisions().get_media(
                            fileId=file_id_final, revisionId=latest_rev['id']
                        )
                    else:
                        req = drive_service.files().get_media(fileId=file_id_final, supportsAllDrives=True)
                except Exception as e_rev:
                    logger.warning(f"Error consultando revisiones de {file_id_final} ({e_rev}), usando get_media estándar.")
                    req = drive_service.files().get_media(fileId=file_id_final, supportsAllDrives=True)
            fh = io.BytesIO()
            dl = MediaIoBaseDownload(fh, req)
            done = False
            while not done:
                _, done = dl.next_chunk()
            return fh.getvalue()
        else:
            nombre_limpio = str(nombre_o_id).strip()
            q_term = nombre_limpio.replace("'", "\\'")
            query = f"name contains '{q_term}' and trashed = false"
            res = drive_service.files().list(
                q=query,
                spaces='drive',
                corpora='allDrives',
                includeItemsFromAllDrives=True,
                supportsAllDrives=True,
                orderBy='modifiedTime desc',
                fields='files(id, name, mimeType)'
            ).execute()
            files = res.get('files', [])

            if not files and '.' in nombre_limpio:
                sin_ext = nombre_limpio.rsplit('.', 1)[0].replace("'", "\\'")
                query2 = f"name contains '{sin_ext}' and trashed = false"
                res2 = drive_service.files().list(
                    q=query2,
                    spaces='drive',
                    corpora='allDrives',
                    includeItemsFromAllDrives=True,
                    supportsAllDrives=True,
                    orderBy='modifiedTime desc',
                    fields='files(id, name, mimeType)'
                ).execute()
                files = res2.get('files', [])

            if files:
                f_id = files[0]['id']
                m_type = files[0].get('mimeType', '')
                if m_type == 'application/vnd.google-apps.document':
                    req = drive_service.files().export_media(fileId=f_id, mimeType='application/pdf')
                else:
                    try:
                        rev_res = drive_service.revisions().list(
                            fileId=f_id, fields='revisions(id, modifiedTime)'
                        ).execute()
                        revs = rev_res.get('revisions', [])
                        if revs:
                            latest_rev = sorted(revs, key=lambda r: r.get('modifiedTime', ''))[-1]
                            logger.info(f"Descargando última revisión {latest_rev.get('id')} ({latest_rev.get('modifiedTime')}) para {f_id}")
                            req = drive_service.revisions().get_media(
                                fileId=f_id, revisionId=latest_rev['id']
                            )
                        else:
                            req = drive_service.files().get_media(fileId=f_id, supportsAllDrives=True)
                    except Exception as e_rev:
                        logger.warning(f"Error consultando revisiones de {f_id} ({e_rev}), usando get_media estándar.")
                        req = drive_service.files().get_media(fileId=f_id, supportsAllDrives=True)
                fh = io.BytesIO()
                dl = MediaIoBaseDownload(fh, req)
                done = False
                while not done:
                    _, done = dl.next_chunk()
                return fh.getvalue()
            else:
                logger.warning(f"No se encontró archivo en Drive con nombre '{nombre_o_id}'")
                return None
    except Exception as e:
        logger.error(f"Error descargando archivo '{nombre_o_id}' desde Drive: {e}")
        return None

# ====================================================================
# --- BLOQUE 5: MANIPULACIÓN DOCUMENTAL (WORD & PDF) ---
# ====================================================================
def _set_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for border_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        b = OxmlElement(f'w:{border_name}')
        b.set(qn('w:val'), 'single')
        b.set(qn('w:sz'), '4')
        b.set(qn('w:space'), '0')
        b.set(qn('w:color'), '000000')
        borders.append(b)
    tblPr.append(borders)

def _set_cell_background(cell, color_hex: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:fill="{color_hex}"/>')
    tcPr.append(shd)

def _set_table_margins(table, top=72, bottom=72, left=30, right=30):
    tblPr = table._tbl.tblPr
    tblCellMar = parse_xml(f'''
    <w:tblCellMar xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
        <w:top w:w="{top}" w:type="dxa"/>
        <w:left w:w="{left}" w:type="dxa"/>
        <w:bottom w:w="{bottom}" w:type="dxa"/>
        <w:right w:w="{right}" w:type="dxa"/>
    </w:tblCellMar>
    ''')
    tblPr.append(tblCellMar)

def ajustar_posicion_y_tamano_firmas(doc: Document, num_items: int) -> None:
    """
    Ajusta las firmas en la plantilla Word (.docx) basándose en la calibración exacta
    del Certificado 219 (1 sola página, espacio limpio entre texto final y firmas,
    sin sobreponerse al pie de página).
    - Firma 1 (Izquierda / EPMI / Gerente Operaciones): cx=1317021 (~3.66 cm), cy=881554 (~2.45 cm), posH=1695450, posV=400050 (~31.5 pt)
    - Firma 2 (Derecha / INECOVE / Resp. Técnico):     cx=1529258 (~4.25 cm), cy=953304 (~2.65 cm), posH=3409950, posV=361187 (~28.4 pt)
    """
    try:
        # 1. Localizar el párrafo de firmas (buscando desde el final hacia arriba)
        sig_p = None
        sig_idx = -1
        total_p = len(doc.paragraphs)
        for i in range(total_p - 1, -1, -1):
            if i < total_p // 2:
                break
            p = doc.paragraphs[i]
            drawings = p._element.findall('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}drawing')
            sig_drawings = []
            for d in drawings:
                ext = d.find('.//{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}extent')
                if ext is not None:
                    try:
                        cy = int(ext.get('cy', 0))
                        # Las firmas tienen alto > 20 pt (254,000 EMUs) y no son marcas de agua de página completa (cy < 4,000,000)
                        if 254000 < cy < 4000000:
                            sig_drawings.append(d)
                    except (ValueError, TypeError):
                        pass
            if len(sig_drawings) >= 1 and ('[[TABLA_NOTAS]]' not in p.text):
                sig_p = p
                sig_idx = i
                break

        if sig_p is None or sig_idx == -1:
            return

        # 2. Eliminar párrafos vacíos posteriores al párrafo de firmas
        paras_posteriores_a_eliminar = []
        for i in range(sig_idx + 1, len(doc.paragraphs)):
            p = doc.paragraphs[i]
            drawings = p._element.findall('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}drawing')
            if not p.text.strip() and not drawings:
                paras_posteriores_a_eliminar.append(p)
        for p in paras_posteriores_a_eliminar:
            try:
                p._element.getparent().remove(p._element)
            except Exception:
                pass

        # 3. Eliminar párrafos vacíos inmediatamente anteriores al párrafo de firmas
        paras_anteriores_a_eliminar = []
        for i in range(sig_idx - 1, -1, -1):
            p = doc.paragraphs[i]
            drawings = p._element.findall('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}drawing')
            if not p.text.strip() and not drawings:
                paras_anteriores_a_eliminar.append(p)
            else:
                break
        for p in paras_anteriores_a_eliminar:
            try:
                p._element.getparent().remove(p._element)
            except Exception:
                pass

        # 4. Ajustar párrafo de texto final inmediatamente anterior ("Se realizó la RECOLECCIÓN...")
        # Calibración Certificado 219: espacio posterior = 0 pt, interlineado = 1.15
        p_prev = None
        for i in range(sig_idx - 1, -1, -1):
            p = doc.paragraphs[i]
            if p.text.strip():
                p_prev = p
                break
        if p_prev is not None:
            p_prev.paragraph_format.space_after = Pt(0)
            p_prev.paragraph_format.line_spacing = 1.15

        # 5. Compactar espaciado del párrafo contenedor de firmas
        # Calibración Certificado 219: space_before = 0 pt, space_after = 0 pt, line_spacing = 1.0
        sig_p.paragraph_format.space_before = Pt(0)
        sig_p.paragraph_format.space_after = Pt(0)
        sig_p.paragraph_format.line_spacing = 1.0

        # Si hay muchos ítems (>= 6), compactar párrafos vacíos intermedios
        if num_items >= 6:
            for p in doc.paragraphs:
                if not p.text.strip() and not p._element.findall('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}drawing'):
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.line_spacing = Pt(2)

        # 6. Calibración exacta de dimensiones y posiciones de firmas basada en Certificado 219:
        # Firma 1 (Izquierda / EPMI): cx=1317021, cy=881554, posH=1695450, posV=400050
        # Firma 2 (Derecha / INECOVE): cx=1529258, cy=953304, posH=3409950, posV=361187
        F1_CX, F1_CY, F1_POS_H, F1_POS_V = 1317021, 881554, 1695450, 400050
        F2_CX, F2_CY, F2_POS_H, F2_POS_V = 1529258, 953304, 3409950, 361187

        # Escalamiento suave solo si la tabla tiene más de 5 ítems para garantizar 1 página
        if num_items <= 5:
            scale = 1.0
        elif num_items == 6:
            scale = 0.95
        else:
            scale = max(0.72, 1.0 - (num_items - 5) * 0.05)

        # 7. Obtener dibujos válidos de firma en el párrafo y ordenarlos de izquierda a derecha (posH)
        drawings = sig_p._element.findall('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}drawing')
        valid_drawings = []
        for d in drawings:
            ext = d.find('.//{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}extent')
            if ext is not None:
                try:
                    cy = int(ext.get('cy', 0))
                    if 254000 < cy < 4000000:
                        valid_drawings.append(d)
                except (ValueError, TypeError):
                    pass

        def _get_pos_h_val(d_elem) -> int:
            posH_elem = d_elem.find('.//{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}positionH')
            if posH_elem is not None:
                posOff = posH_elem.find('{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}posOffset')
                if posOff is not None:
                    try:
                        return int(posOff.text)
                    except (ValueError, TypeError):
                        pass
            return 0

        valid_drawings.sort(key=_get_pos_h_val)

        specs = [
            {'cx': int(F1_CX * scale), 'cy': int(F1_CY * scale), 'posH': F1_POS_H, 'posV': int(F1_POS_V * scale)},
            {'cx': int(F2_CX * scale), 'cy': int(F2_CY * scale), 'posH': F2_POS_H, 'posV': int(F2_POS_V * scale)}
        ]

        for idx_d, d in enumerate(valid_drawings):
            spec = specs[idx_d] if idx_d < len(specs) else specs[-1]

            # Eliminar márgenes de envoltura en anchor
            for anchor in d.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}anchor'):
                anchor.set('distT', '0')
                anchor.set('distB', '0')
                anchor.set('distL', '114300')
                anchor.set('distR', '114300')

            # Posición horizontal fija calibrada
            posH = d.find('.//{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}positionH')
            if posH is not None:
                posOffsetH = posH.find('{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}posOffset')
                if posOffsetH is not None:
                    posOffsetH.text = str(spec['posH'])

            # Posición vertical fija respecto al párrafo
            posV = d.find('.//{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}positionV')
            if posV is not None:
                posOffsetV = posV.find('{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}posOffset')
                if posOffsetV is not None:
                    posOffsetV.text = str(spec['posV'])

            # Ajustar dimensiones exactas en todos los elementos con cx y cy
            for el in d.iter():
                if 'cx' in el.attrib and 'cy' in el.attrib:
                    el.attrib['cx'] = str(spec['cx'])
                    el.attrib['cy'] = str(spec['cy'])

    except Exception as e:
        logger.warning(f"Aviso en ajuste de firmas: {e}")

def inyectar_tabla_en_docx(doc_io: io.BytesIO, data_items: List[Dict[str, Any]], columnas: Optional[List[Dict[str, Any]]] = None) -> bytes:
    """
    Inyecta la tabla formateada en el marcador [[TABLA_NOTAS]] del documento Word.
    Soporta columnas dinámicas (encabezados editables y eliminables).
    """
    doc = Document(doc_io)
    target_paragraph = None
    for p in doc.paragraphs:
        if '[[TABLA_NOTAS]]' in p.text:
            target_paragraph = p
            break

    if target_paragraph:
        target_paragraph.text = target_paragraph.text.replace('[[TABLA_NOTAS]]', '')
        for p in doc.paragraphs[:5]:
            if "CERTIFICADO" in p.text.upper():
                p.paragraph_format.space_after = Pt(0)

        # Resolver columnas activas y sus encabezados
        if isinstance(columnas, str):
            try:
                columnas = json.loads(columnas)
            except Exception:
                columnas = None

        if not columnas or not isinstance(columnas, list) or len(columnas) == 0:
            columnas_def = [
                {"key": "fecha_origen", "label": "Fecha"},
                {"key": "placa_origen", "label": "Placa"},
                {"key": "guia_origen", "label": "N° Guía"},
                {"key": "desc", "label": "Descripción"},
                {"key": "cant", "label": "Cantidad"},
                {"key": "um", "label": "Medida"},
                {"key": "peso", "label": "Peso (Kg)"}
            ]
        else:
            columnas_def = []
            for c_item in columnas:
                if isinstance(c_item, dict) and c_item.get('key'):
                    columnas_def.append({
                        "key": str(c_item.get('key')).strip(),
                        "label": str(c_item.get('label') or c_item.get('key')).strip()
                    })
            if not columnas_def:
                columnas_def = [
                    {"key": "fecha_origen", "label": "Fecha"},
                    {"key": "placa_origen", "label": "Placa"},
                    {"key": "guia_origen", "label": "N° Guía"},
                    {"key": "desc", "label": "Descripción"},
                    {"key": "cant", "label": "Cantidad"},
                    {"key": "um", "label": "Medida"},
                    {"key": "peso", "label": "Peso (Kg)"}
                ]

        num_cols = len(columnas_def)
        table = doc.add_table(rows=1, cols=num_cols)
        try:
            table.style = 'Table Grid'
        except Exception:
            _set_table_borders(table)

        table.autofit = False
        table.allow_autofit = False
        _set_table_margins(table, top=72, bottom=72, left=30, right=30)

        # Distribuir anchos de columnas dinámicamente para ocupar el ancho imprimible (7.5 pulgadas)
        has_desc = any(c_item.get('key') == 'desc' for c_item in columnas_def)
        widths = []
        if has_desc and num_cols > 1:
            other_w = min(1.0, max(0.65, (7.5 - 3.0) / (num_cols - 1)))
            desc_w = 7.5 - (other_w * (num_cols - 1))
            for c_item in columnas_def:
                if c_item.get('key') == 'desc':
                    widths.append(Inches(desc_w))
                else:
                    widths.append(Inches(other_w))
        else:
            equal_w = 7.5 / num_cols
            for _ in columnas_def:
                widths.append(Inches(equal_w))

        for i, col in enumerate(table.columns):
            col.width = widths[i]

        hdr_cells = table.rows[0].cells
        for i, col_info in enumerate(columnas_def):
            nombre = col_info['label']
            cell = hdr_cells[i]
            cell.text = nombre
            cell.width = widths[i]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            _set_cell_background(cell, "70ad47")

            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1
                run = p.runs[0] if p.runs else p.add_run(nombre)
                run.font.bold = True
                run.font.color.rgb = RGBColor(0, 0, 0)
                run.font.name = 'Calibri'
                run.font.size = Pt(9)

        for item in data_items:
            row_cells = table.add_row().cells
            vals = []
            for col_info in columnas_def:
                k = col_info['key']
                if k == 'guia_origen':
                    v = limpiar_numero_guia(item.get('guia_origen', ''))
                elif k == 'um':
                    v = str(item.get('um', '')).upper()
                else:
                    v = str(item.get(k, ''))
                vals.append(v)

            for idx, valor in enumerate(vals):
                cell = row_cells[idx]
                cell.text = valor
                cell.width = widths[idx]
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.line_spacing = 1
                    run = p.runs[0] if p.runs else p.add_run(valor)
                    run.font.name = 'Calibri'
                    run.font.size = Pt(9)

        tbl, p_elem = table._tbl, target_paragraph._p
        p_elem.addnext(tbl)

    # Ajuste adaptativo de firmas para evitar superposición con el pie de página
    ajustar_posicion_y_tamano_firmas(doc, len(data_items))

    new_buffer = io.BytesIO()
    doc.save(new_buffer)
    docx_bytes_result = new_buffer.getvalue()
    new_buffer.close()
    del doc
    return docx_bytes_result

def convertir_docx_a_pdf(docx_bytes: bytes, drive_service=None) -> bytes:
    """
    Convierte un documento .docx a PDF usando la API de Google Drive
    (sin requerir Microsoft Word ni LibreOffice en el host).
    """
    if not drive_service:
        drive_service, _ = obtener_servicios_google()

    upload_io = io.BytesIO(docx_bytes)
    try:
        media = MediaIoBaseUpload(
            upload_io,
            mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
            resumable=False
        )
        temp_file = drive_service.files().create(
            body={'name': f"temp_conv_{int(time.time())}", 'mimeType': 'application/vnd.google-apps.document'},
            media_body=media,
            fields='id'
        ).execute()
        temp_id = temp_file.get('id')

        req = drive_service.files().export_media(fileId=temp_id, mimeType='application/pdf')
        fh = io.BytesIO()
        dl = MediaIoBaseDownload(fh, req)
        done = False
        while not done:
            _, done = dl.next_chunk()

        # Limpiar archivo temporal en Drive
        try:
            drive_service.files().delete(fileId=temp_id).execute()
        except Exception:
            pass

        pdf_result = fh.getvalue()
        fh.close()
        return pdf_result
    except Exception as e:
        logger.error(f"Error convirtiendo DOCX a PDF vía Drive API: {e}")
        raise RuntimeError(f"Falla convirtiendo Word a PDF: {e}")
    finally:
        upload_io.close()

def unir_tres_documentos_pdf(cert_pdf_bytes: bytes, remision_bytes: Optional[bytes] = None, transporte_bytes: Optional[bytes] = None) -> bytes:
    """
    Une en orden estricto los documentos:
    1. Certificado (Página inicial)
    2. Guía de Remisión (Páginas posteriores)
    3. Guía de Transporte (Opcional)
    """
    writer = PdfWriter()
    io_cert = io.BytesIO(cert_pdf_bytes)
    reader_cert = PdfReader(io_cert)
    for page in reader_cert.pages:
        writer.add_page(page)

    io_rem = None
    if remision_bytes:
        try:
            io_rem = io.BytesIO(remision_bytes)
            reader_rem = PdfReader(io_rem)
            for page in reader_rem.pages:
                writer.add_page(page)
        except Exception as e:
            logger.warning(f"No se pudo unir remisión como PDF: {e}")

    io_trans = None
    if transporte_bytes:
        try:
            io_trans = io.BytesIO(transporte_bytes)
            reader_trans = PdfReader(io_trans)
            for page in reader_trans.pages:
                writer.add_page(page)
        except Exception as e:
            logger.warning(f"No se pudo unir transporte como PDF: {e}")

    output = io.BytesIO()
    writer.write(output)
    result = output.getvalue()

    # Cerrar buffers y liberar memoria explícitamente
    output.close()
    io_cert.close()
    if io_rem:
        io_rem.close()
    if io_trans:
        io_trans.close()
    del writer, reader_cert
    liberar_memoria_proceso()
    return result

def sustituir_certificado_en_pdf(pdf_unificado_bytes: bytes, nuevo_cert_pdf_bytes: bytes, num_paginas_reemplazar: int = 1) -> bytes:
    """
    Sustituye quirúrgicamente la(s) primera(s) página(s) de un PDF consolidado
    conservando intactas todas las páginas de guías originales posteriores.
    """
    io_unido = io.BytesIO(pdf_unificado_bytes)
    io_nuevo = io.BytesIO(nuevo_cert_pdf_bytes)
    reader_unido = PdfReader(io_unido)
    reader_nuevo = PdfReader(io_nuevo)
    writer = PdfWriter()

    # 1. Agregar páginas del nuevo certificado
    for p in reader_nuevo.pages:
        writer.add_page(p)

    # 2. Agregar las páginas restantes del PDF unificado original (las guías)
    total_unido = len(reader_unido.pages)
    paginas_a_saltar = min(num_paginas_reemplazar, total_unido)
    for p in reader_unido.pages[paginas_a_saltar:]:
        writer.add_page(p)

    out_io = io.BytesIO()
    writer.write(out_io)
    result = out_io.getvalue()

    out_io.close()
    io_unido.close()
    io_nuevo.close()
    del writer, reader_unido, reader_nuevo
    liberar_memoria_proceso()
    return result

# ====================================================================
# --- BLOQUE 6: MOTOR COGNITIVO IA (VERTEX AI / GEMINI OCR) ---
# ====================================================================
PROMPT_EXTRACCION_CERTIFICADO = """
INSTRUCCIÓN DE SISTEMA: Eres un extractor de datos OCR estricto. Tu única tarea es extraer datos del PDF adjunto y devolverlos ÚNICAMENTE en formato JSON válido. Tienes PROHIBIDO inventar datos, alucinar información o incluir texto fuera del JSON (como ```json o explicaciones).

ESTRUCTURA JSON EXACTA Y REGLAS DE NEGOCIO OBLIGATORIAS:
{
    "cliente": "Razón Social exacta del REMITENTE. Regla Estricta: NO extraer la empresa de transportes, NO extraer nombres de conductores.",
    "ruc_cliente": "Número de RUC del Remitente o Cliente Emisor.",
    "fecha": "dd/mm/yyyy", 
    "serie": "Serie-Numero completo de la guía. Ejemplo: T001-000000", 
    "vehiculo": "PLACA del vehículo. Busca en todo el documento. Obligatorio.", 
    
    "punto_partida": "REGLA DE ORO OBLIGATORIA: Lee primero el bloque 'Observaciones' u 'Observación' de la guía. Todo dato como 'Fundo Casuarinas', 'Planta...', u otro predio que aparezca ahí TIENE QUE SER EXTRAÍDO SÍ O SÍ. Concatena la dirección base de partida con ese dato usando un guion. Ejemplo de Salida Exacta: 'Direccion Base - Fundo Casuarinas' o 'Av Sur - PLANTA EMPACADORA'. Si dice textualmente 'Fundo Casuarinas', debe salir 'Fundo Casuarinas'. NUNCA dejes fuera la información de 'Observaciones'. Si este campo está vacío entonces devuelve solo la dirección de partida base. NUNCA deduzcas ni inventes basándote en la empresa.", 
    
    "punto_llegada": "Dirección Completa exacta de Llegada. IMPORTANTE: Si en el documento (especialmente para la empresa Los Olivos de Villacuri) el destino o planta se indica simplemente como 'EMPACADORA', debes extraer la palabra 'EMPACADORA' y asignarla obligatoriamente a este campo. No lo dejes vacío.", 
    "destinatario": "Razón Social Completa del Destinatario", 
    
    "items": [
        {
            "desc": "Descripción literal del bien", 
            "cant": "Número", 
            "um": "Unidad de medida (KG, UNID, GLN)", 
            "peso": "Peso numérico explícito (o 0.00 si no existe)"
        }
    ]
}
"""

_cached_genai_clients: Dict[str, Any] = {}

def procesar_guia_ia_vertex(pdf_bytes: bytes) -> Optional[Dict[str, Any]]:
    """Procesamiento OCR síncrono con Gemini vía Vertex AI / Google GenAI SDK con clientes cacheados y liberación de memoria."""
    global _cached_genai_clients
    creds = obtener_credenciales()
    regiones = [REGION_ESTABLE or "us-central1", "us-west1", "us-east4", "southamerica-east1"]
    modelos = [MODEL_NAME or "gemini-2.5-flash"] + FALLBACK_MODELS

    pdf_part = types.Part.from_bytes(data=pdf_bytes, mime_type="application/pdf")
    errores = []

    try:
        for region in regiones:
            try:
                if region not in _cached_genai_clients:
                    _cached_genai_clients[region] = genai.Client(vertexai=True, project=PROJECT_ID, location=region, credentials=creds)
                client = _cached_genai_clients[region]

                for m_name in modelos:
                    try:
                        response = client.models.generate_content(
                            model=m_name,
                            contents=[pdf_part, PROMPT_EXTRACCION_CERTIFICADO],
                            config=types.GenerateContentConfig(response_mime_type="application/json")
                        )
                        datos = json.loads(response.text)
                        if datos.get("destinatario") or len(datos.get("vehiculo", "")) >= 3 or datos.get("cliente"):
                            if datos.get("serie"):
                                datos["serie"] = limpiar_numero_guia(datos["serie"])
                            return datos
                    except Exception as e:
                        if "404" not in str(e):
                            errores.append(f"{region}/{m_name}: {e}")
                        continue
            except Exception as e:
                errores.append(f"Init {region}: {e}")
                continue

        logger.error(f"Falla OCR de Vertex AI en todas las regiones: {errores}")
        return None
    finally:
        del pdf_part
        liberar_memoria_proceso()

async def async_procesar_guia_ia_vertex(pdf_bytes: bytes) -> Optional[Dict[str, Any]]:
    """Envoltorio asíncrono para no bloquear el bucle de eventos del Bot."""
    return await asyncio.to_thread(procesar_guia_ia_vertex, pdf_bytes)

# ====================================================================
# --- BLOQUE 7: LECTURA DE CATÁLOGOS Y CORRELATIVOS ---
# ====================================================================
def obtener_catalogo_empresas(sheets_service=None) -> Dict[str, Dict[str, str]]:
    """Lee la pestaña 'EMPRESAS' (A=Nombre, B=RUC, C=Registro)."""
    if not sheets_service:
        _, sheets_service = obtener_servicios_google()
    try:
        res = sheets_service.spreadsheets().values().get(
            spreadsheetId=ID_SHEET_CONTROL, range="EMPRESAS!A2:C"
        ).execute()
        filas = res.get('values', [])
        catalogo = {}
        for fila in filas:
            if len(fila) >= 2:
                nombre = str(fila[0]).strip().upper()
                ruc = str(fila[1]).strip()
                reg = str(fila[2]).strip() if len(fila) >= 3 else "Pendiente"
                if nombre:
                    catalogo[nombre] = {"ruc": ruc, "reg": reg}
        return catalogo
    except Exception as e:
        logger.error(f"Error leyendo catálogo EMPRESAS: {e}")
        return {
            "EPMI S.A.C.": {"ruc": "20601369792", "reg": "EO-RS-00043-2021-MINAM"},
            "INECOVE S.A.C.": {"ruc": "20607730101", "reg": "EO-RS-00109-2022-MINAM"}
        }

def obtener_catalogo_servicios(sheets_service=None) -> Dict[str, Dict[str, List[str]]]:
    """Lee y agrupa la pestaña 'SERVICIOS' (Comercialización vs Servicios)."""
    if not sheets_service:
        _, sheets_service = obtener_servicios_google()
    secciones = {
        "COMERCIALIZACION": {"titulos": [], "servicios": [], "residuos": []},
        "SERVICIOS": {"titulos": [], "servicios": [], "residuos": []}
    }
    try:
        res = sheets_service.spreadsheets().values().get(
            spreadsheetId=ID_SHEET_CONTROL, range="SERVICIOS!A1:C500"
        ).execute()
        filas = res.get('values', [])
        if not filas:
            return secciones

        seccion_actual = "COMERCIALIZACION"
        KEYWORDS_HEADER = ["COMERCIALIZACION", "COMERCIALIZACIÓN", "SERVICIOS", "DISPOSICION FINAL", "DISPOSICIÓN FINAL"]
        norm_keywords = [normalizar_texto_sin_tildes(k).upper() for k in KEYWORDS_HEADER]

        for row in filas:
            c0 = str(row[0]).strip() if len(row) > 0 else ""
            c1 = str(row[1]).strip() if len(row) > 1 else ""
            c2 = str(row[2]).strip() if len(row) > 2 else ""

            c0_clean = " ".join(c0.split())
            c1_clean = " ".join(c1.split())
            c2_clean = " ".join(c2.split())

            c0_norm = normalizar_texto_sin_tildes(c0_clean).upper()

            # Ignorar encabezados de tabla (fila 1: TITULOS, SERVICIO O COMPRA, TIPO DE RESIDUO)
            if c0_norm in ["TITULOS", "TITULO"]:
                continue

            # Detectar cambio de sección por coincidencia EXACTA
            if c0_norm in norm_keywords:
                seccion_actual = "COMERCIALIZACION" if "COMERC" in c0_norm else "SERVICIOS"
                continue

            if seccion_actual in secciones:
                if c0:
                    tipo_sec = "COM" if seccion_actual == "COMERCIALIZACION" else "SER"
                    c0_fmt = formatear_titulo_oficial(c0, tipo_cod=tipo_sec)
                    if c0_fmt and c0_fmt not in secciones[seccion_actual]["titulos"]:
                        secciones[seccion_actual]["titulos"].append(c0_fmt)
                if c1_clean and c1_clean not in secciones[seccion_actual]["servicios"]:
                    secciones[seccion_actual]["servicios"].append(c1_clean)
                if c2_clean and c2_clean not in secciones[seccion_actual]["residuos"]:
                    secciones[seccion_actual]["residuos"].append(c2_clean)

        for sec, default_t in [
            ("COMERCIALIZACION", "CERTIFICADO DE RECOLECCIÓN, TRANSPORTE Y \nCOMERCIALIZACIÓN DE RESIDUOS SÓLIDOS"),
            ("SERVICIOS", "CERTIFICADO DE RECOLECCIÓN Y TRANSPORTE \nDE RESIDUOS SÓLIDOS PELIGROSOS")
        ]:
            if not secciones[sec]["titulos"]:
                secciones[sec]["titulos"].append(default_t)

        return secciones
    except Exception as e:
        logger.error(f"Error leyendo pestaña SERVICIOS: {e}")
        return secciones

def obtener_catalogo_clientes_certificados(sheets_service=None) -> Dict[str, Dict[str, str]]:
    """Lee la pestaña 'CLIENTES' (A=Empresa, B=RUC, C=Registro, D=Dirección Certificado, E=Domicilio Fiscal)."""
    if not sheets_service:
        _, sheets_service = obtener_servicios_google()
    try:
        res = sheets_service.spreadsheets().values().get(
            spreadsheetId=ID_SHEET_CONTROL, range="CLIENTES!A2:E"
        ).execute()
        filas = res.get('values', [])
        catalogo = {}
        for fila in filas:
            if len(fila) >= 1:
                empresa = str(fila[0]).strip().upper()
                if not empresa or empresa in ["S/D", "EMPRESA"]:
                    continue
                ruc = str(fila[1]).strip() if len(fila) > 1 else ""
                registro = str(fila[2]).strip() if len(fila) > 2 else ""
                dir_cert = str(fila[3]).strip() if len(fila) > 3 else ""
                dom_fisc = str(fila[4]).strip() if len(fila) > 4 else ""
                catalogo[empresa] = {
                    "empresa": empresa,
                    "ruc": ruc,
                    "registro": registro,
                    "direccion": dir_cert or dom_fisc
                }
        return catalogo
    except Exception as e:
        logger.error(f"Error leyendo catálogo CLIENTES: {e}")
        return {}

def obtener_catalogo_direcciones_certificados(sheets_service=None) -> Dict[str, Dict[str, str]]:
    """Lee la pestaña 'Direcciones' (A=Empresa, B=Fundo/Planta, C=Palabra Clave, D=Dirección)."""
    if not sheets_service:
        _, sheets_service = obtener_servicios_google()
    try:
        res = sheets_service.spreadsheets().values().get(
            spreadsheetId=ID_SHEET_CONTROL, range="Direcciones!A2:D"
        ).execute()
        filas = res.get('values', [])
        catalogo = {}
        for fila in filas:
            if len(fila) >= 4:
                empresa = str(fila[0]).strip().upper()
                fundo = str(fila[1]).strip().upper()
                direccion = str(fila[3]).strip()
                if empresa and fundo and direccion:
                    if empresa not in catalogo:
                        catalogo[empresa] = {}
                    catalogo[empresa][fundo] = direccion
        return catalogo
    except Exception as e:
        logger.error(f"Error leyendo catálogo Direcciones: {e}")
        return {}

def obtener_siguiente_correlativo_cert(tipo_flujo: str = "Comercialización", sheets_service=None) -> str:
    """Calcula el siguiente número correlativo sugerido leyendo 'historial'."""
    if not sheets_service:
        _, sheets_service = obtener_servicios_google()
    try:
        res = sheets_service.spreadsheets().values().get(
            spreadsheetId=ID_SHEET_CONTROL, range="'historial'!D2:E"
        ).execute()
        filas = res.get('values', [])
        max_num = 0
        tipo_filtro = "COM" if "comercializa" in str(tipo_flujo).lower() else "SER"

        for row in filas:
            corr_val = str(row[0]).strip() if len(row) > 0 else ""
            tipo_val = str(row[1]).strip() if len(row) > 1 else ""

            # Si el tipo coincide o no está especificado
            es_com = "comercializa" in tipo_val.lower() or "com" in tipo_val.lower()
            if (tipo_filtro == "COM" and es_com) or (tipo_filtro != "COM" and not es_com):
                nums = re.findall(r'\d+', corr_val)
                if nums:
                    num_int = int(nums[-1])
                    if num_int > max_num:
                        max_num = num_int

        sig = max_num + 1 if max_num > 0 else 1
        return f"{sig:03d}"
    except Exception as e:
        logger.error(f"Error calculando siguiente correlativo: {e}")
        return "001"

# ====================================================================
# --- BLOQUE 8: GESTIÓN DE PLANTILLAS Y SUBIDA A DRIVE ---
# ====================================================================
def resolver_plantilla_drive(empresa_nombre: str, tipo_certificado: str, drive_service, es_modelo: bool = False) -> str:
    """
    Busca la plantilla .docx en la carpeta de Drive o retorna el ID fallback.
    Garantiza que 'Comercialización' nunca use plantillas de Disposición Final.
    """
    carpeta_id = CARPETA_PLANTILLAS_MODELO if es_modelo else CARPETA_PLANTILLAS_NORMALES
    emp_key = "INECOVE S.A.C." if "INECOVE" in empresa_nombre.upper() else "EPMI S.A.C."
    empresa_limpia = "INECOVE" if "INECOVE" in empresa_nombre.upper() else "EPMI"

    tipo_str = str(tipo_certificado or "").lower()
    if "comercial" in tipo_str:
        flujo_key = "Comercialización"
        palabra_flujo = "COMERCIALIZACION"
    elif "2" in tipo_str:
        flujo_key = "Disposición Final 2"
        palabra_flujo = "final 2"
    else:
        flujo_key = "Disposición Final 1"
        palabra_flujo = "final 1"

    if es_modelo:
        palabra_modelo = "COMERCIALIZACION" if flujo_key == "Comercialización" else "FINAL"
        try:
            query = f"'{carpeta_id}' in parents and name contains '{palabra_modelo}' and trashed = false"
            res = drive_service.files().list(q=query, spaces='drive', fields='files(id, name)', supportsAllDrives=True).execute()
            files = res.get('files', [])
            if files:
                return files[0]['id']
        except Exception as e:
            logger.warning(f"Búsqueda dinámica de plantilla modelo falló ({e}).")
    else:
        try:
            query = (
                f"'{carpeta_id}' in parents "
                f"and name contains '{empresa_limpia}' "
                f"and name contains '{palabra_flujo}' "
                f"and trashed = false"
            )
            res = drive_service.files().list(
                q=query, spaces='drive', fields='files(id, name, mimeType)', supportsAllDrives=True
            ).execute()
            files = res.get('files', [])
            if files:
                return files[0]['id']
        except Exception as e:
            logger.warning(f"Búsqueda dinámica de plantilla falló ({e}). Usando fallback.")

    # Fallback predeterminado según empresa y flujo
    return PLANTILLAS_FALLBACK.get(emp_key, {}).get(flujo_key, "1d09vmlBlW_4yjrrz5M1XM8WpCvzTI4f11pERDbxFvNE")

def descargar_plantilla_docx(plantilla_id: str, drive_service) -> bytes:
    """Descarga los bytes de una plantilla de Drive (Doc o Docx) asegurando la última revisión."""
    meta = drive_service.files().get(fileId=plantilla_id, fields='mimeType', supportsAllDrives=True).execute()
    mime = meta.get('mimeType', '')
    if mime == 'application/vnd.google-apps.document':
        req = drive_service.files().export_media(
            fileId=plantilla_id,
            mimeType='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        )
    else:
        # Para archivos binarios (.docx), consultar revisiones para descargar siempre la última revisión real
        try:
            rev_res = drive_service.revisions().list(
                fileId=plantilla_id, fields='revisions(id, modifiedTime)'
            ).execute()
            revs = rev_res.get('revisions', [])
            if revs:
                latest_rev = sorted(revs, key=lambda r: r.get('modifiedTime', ''))[-1]
                logger.info(f"Descargando última revisión de Drive {latest_rev.get('id')} ({latest_rev.get('modifiedTime')}) para {plantilla_id}")
                req = drive_service.revisions().get_media(
                    fileId=plantilla_id, revisionId=latest_rev['id']
                )
            else:
                req = drive_service.files().get_media(fileId=plantilla_id, supportsAllDrives=True)
        except Exception as e_rev:
            logger.warning(f"No se pudo consultar revisiones de {plantilla_id} ({e_rev}), usando get_media estándar.")
            req = drive_service.files().get_media(fileId=plantilla_id, supportsAllDrives=True)

    fh = io.BytesIO()
    dl = MediaIoBaseDownload(fh, req)
    done = False
    while not done:
        _, done = dl.next_chunk()
    return fh.getvalue()

def subir_a_drive_certificado(contenido_bytes: bytes, nombre_archivo: str, tipo_flujo: str, empresa_firma: str, es_modelo: bool = False, mime_type: str = 'application/pdf') -> Tuple[Optional[str], Optional[str]]:
    """
    Sube el archivo generado a Google Drive en la carpeta de destino correspondiente.
    Retorna (file_id, webViewLink).
    """
    drive_service, _ = obtener_servicios_google()
    if es_modelo:
        carpeta_id = "1LUErbILxjVHnzuHkdWaeAMI4HnLg1c7E"
    else:
        emp_key = "INECOVE S.A.C." if "INECOVE" in empresa_firma.upper() else "EPMI S.A.C."
        flujo_key = tipo_flujo if tipo_flujo in CARPETAS_DESTINO.get(emp_key, {}) else "Comercialización"
        carpeta_id = CARPETAS_DESTINO.get(emp_key, {}).get(flujo_key, "1NZc-nfGHw5bnkCAv0TdQYW_bPM_UkKC-")

    try:
        media = MediaIoBaseUpload(io.BytesIO(contenido_bytes), mimetype=mime_type, resumable=True)
        file_metadata = {
            'name': nombre_archivo,
            'parents': [carpeta_id]
        }
        res = drive_service.files().create(
            body=file_metadata, media_body=media, fields='id, webViewLink', supportsAllDrives=True
        ).execute()
        return res.get('id'), res.get('webViewLink')
    except Exception as e:
        logger.error(f"Error subiendo archivo a Drive: {e}")
        return None, None

def sobrescribir_o_subir_pdf_drive(file_id_existente: Optional[str], contenido_bytes: bytes, nombre_archivo: str, tipo_flujo: str, empresa_firma: str) -> Optional[str]:
    """
    Actualiza in-place el archivo PDF consolidado en Google Drive para preservar su ID y URL pública.
    """
    drive_service, _ = obtener_servicios_google()
    if nombre_archivo:
        nombre_archivo = nombre_archivo.strip()
        if not nombre_archivo.lower().endswith('.pdf') and not nombre_archivo.startswith(('http://', 'https://')):
            nombre_archivo += '.pdf'

    if file_id_existente:
        try:
            media = MediaIoBaseUpload(io.BytesIO(contenido_bytes), mimetype='application/pdf', resumable=True)
            body = {'name': nombre_archivo} if nombre_archivo else {}
            res = drive_service.files().update(
                fileId=file_id_existente,
                body=body,
                media_body=media,
                fields='id, name, webViewLink',
                supportsAllDrives=True
            ).execute()
            link = res.get('webViewLink')
            if link:
                return link
        except Exception as e:
            logger.warning(f"Falla actualización in-place en Drive ({e}). Subiendo nuevo archivo...")

    _, link_nuevo = subir_a_drive_certificado(contenido_bytes, nombre_archivo, tipo_flujo, empresa_firma, mime_type='application/pdf')
    return link_nuevo

def buscar_fila_certificado_existente(correlativo: str, tipo_cod: str = "COM", sheets_service=None) -> Optional[int]:
    """
    Busca si ya existe una fila para el correlativo y tipo especificado en la pestaña 'historial'.
    Retorna el número de fila (1-based) o None si no existe.
    Escanea desde las filas más recientes hacia atrás.
    """
    if not correlativo:
        return None
    if not sheets_service:
        _, sheets_service = obtener_servicios_google()
    try:
        res = sheets_service.spreadsheets().values().get(
            spreadsheetId=ID_SHEET_CONTROL, range="'historial'!A2:E"
        ).execute()
        filas = res.get('values', [])
        if not filas:
            return None

        corr_query = str(correlativo).strip().upper()
        corr_num = int(corr_query) if corr_query.isdigit() else None
        tipo_query = str(tipo_cod or "").strip().upper()

        # Recorrer de abajo hacia arriba (los más recientes primero)
        for idx in range(len(filas) - 1, -1, -1):
            row = filas[idx]
            num_fila = idx + 2
            corr_row = str(row[3]).strip().upper() if len(row) > 3 else ""
            if not corr_row:
                continue

            match_corr = False
            if corr_row == corr_query:
                match_corr = True
            elif corr_num is not None and corr_row.isdigit() and int(corr_row) == corr_num:
                match_corr = True

            if match_corr:
                tipo_row = str(row[4]).strip().upper() if len(row) > 4 else ""
                if not tipo_query:
                    return num_fila

                if tipo_query == "COM":
                    if "COM" in tipo_row:
                        return num_fila
                else:  # "SER"
                    if any(k in tipo_row for k in ["FIN", "SER", "DISPOSIC"]):
                        return num_fila

        return None
    except Exception as e:
        logger.error(f"Error buscando fila existente en historial: {e}")
        return None

def actualizar_fila_historial_sheets(num_fila: int, datos_fila: List[Any], sheets_service=None) -> bool:
    """
    Actualiza completamente una fila existente en la pestaña 'historial' (A{num_fila}:J{num_fila}).
    Sobrescribe todos los campos incluyendo la auditoría en la columna J.
    """
    if not num_fila or num_fila < 2:
        return False
    if not sheets_service:
        _, sheets_service = obtener_servicios_google()
    try:
        sheets_service.spreadsheets().values().update(
            spreadsheetId=ID_SHEET_CONTROL,
            range=f"'historial'!A{num_fila}:J{num_fila}",
            valueInputOption="USER_ENTERED",
            body={"values": [datos_fila]}
        ).execute()
        return True
    except Exception as e:
        logger.error(f"Error actualizando fila {num_fila} en 'historial': {e}")
        return False

def registrar_en_historial_sheets(datos_fila: List[Any], sheets_service=None) -> Optional[int]:
    """Inserta una fila en la pestaña 'historial' (A:J)."""
    if not sheets_service:
        _, sheets_service = obtener_servicios_google()
    try:
        res = sheets_service.spreadsheets().values().append(
            spreadsheetId=ID_SHEET_CONTROL,
            range="'historial'!A:J",
            valueInputOption="USER_ENTERED",
            insertDataOption="INSERT_ROWS",
            body={"values": [datos_fila]}
        ).execute()
        updates = res.get('updates', {})
        updated_range = updates.get('updatedRange', '')
        m = re.search(r'(\d+)', updated_range.split('!')[-1])
        return int(m.group(1)) if m else None
    except Exception as e:
        logger.error(f"Error registrando en 'historial': {e}")
        return None

def registrar_edicion_en_historial(num_fila: int, link_pdf: str, usuario_editor: str = "Usuario", obs_extra: str = "", sheets_service=None) -> bool:
    """Actualiza columnas I y J de 'historial' tras una edición quirúrgica o unificación."""
    if not sheets_service:
        _, sheets_service = obtener_servicios_google()
    if not num_fila or num_fila < 2:
        return False
    try:
        ahora_pe = (datetime.now(PET)).strftime("%d/%m/%Y %H:%M")
        u_clean = str(usuario_editor or "Usuario").strip()
        if u_clean.startswith('@'):
            u_clean = u_clean[1:]
        nota_audit = f"Modificado por Mini App ({ahora_pe} - {u_clean})"
        if obs_extra:
            nota_audit += f" | {obs_extra.strip()}"
        body = {"values": [[link_pdf, nota_audit]]}
        sheets_service.spreadsheets().values().update(
            spreadsheetId=ID_SHEET_CONTROL,
            range=f"'historial'!I{num_fila}:J{num_fila}",
            valueInputOption="USER_ENTERED",
            body=body
        ).execute()
        return True
    except Exception as e:
        logger.error(f"Error registrando edición en historial fila {num_fila}: {e}")
        return False

# ====================================================================
# --- BLOQUE 9: BÚSQUEDA DE CERTIFICADOS PARA MODIFICACIÓN ---
# ====================================================================
def buscar_datos_certificado_en_historial(correlativo: str, sheets_service=None, drive_service=None) -> List[Dict[str, Any]]:
    """
    Busca certificados en 'historial' por número correlativo (ej. '045' o '45').
    Resuelve los enlaces de Google Docs y PDF en Drive.
    """
    if not sheets_service or not drive_service:
        drive_service, sheets_service = obtener_servicios_google()

    if not correlativo:
        return []

    try:
        res = sheets_service.spreadsheets().values().get(
            spreadsheetId=ID_SHEET_CONTROL, range="'historial'!A:J",
            valueRenderOption="FORMULA"
        ).execute()
        filas = res.get('values', [])
        if len(filas) < 2:
            return []

        corr_query = str(correlativo).strip().upper()
        corr_num = int(corr_query) if corr_query.isdigit() else None
        resultados_exactos = []
        resultados_parciales = []

        # Recorrer de abajo hacia arriba (más recientes primero)
        for idx in range(len(filas) - 1, 0, -1):
            row = filas[idx]
            num_fila = idx + 1
            corr_row = str(row[3]).strip().upper().replace('=', '').replace('"', '') if len(row) > 3 else ""
            if not corr_row:
                continue

            es_exacto = False
            es_parcial = False
            if corr_row == corr_query:
                es_exacto = True
            elif corr_num is not None and corr_row.isdigit() and int(corr_row) == corr_num:
                es_exacto = True
            elif corr_query and corr_query in corr_row:
                es_parcial = True

            if es_exacto or es_parcial:
                raw_guia = str(row[6]).strip() if len(row) > 6 else ""
                raw_doc_val = str(row[7]).strip() if len(row) > 7 else ""
                raw_pdf_val = str(row[8]).strip() if len(row) > 8 else ""

                url_doc, nom_doc = parsear_celda_enlace_o_nombre(raw_doc_val)
                url_pdf, nom_pdf = parsear_celda_enlace_o_nombre(raw_pdf_val)

                # Resolver enlace y nombre para DOCX
                if not url_doc and nom_doc:
                    url_doc = obtener_link_archivo_drive(drive_service, nom_doc) or ""
                elif url_doc and not nom_doc:
                    doc_id = extraer_id_drive(url_doc)
                    if doc_id:
                        try:
                            meta = drive_service.files().get(fileId=doc_id, fields='name', supportsAllDrives=True).execute()
                            nom_doc = meta.get('name', '')
                        except Exception:
                            pass

                # Resolver enlace y nombre para PDF
                if not url_pdf and nom_pdf:
                    url_pdf = obtener_link_archivo_drive(drive_service, nom_pdf) or ""
                elif url_pdf and not nom_pdf:
                    pdf_id = extraer_id_drive(url_pdf)
                    if pdf_id:
                        try:
                            meta = drive_service.files().get(fileId=pdf_id, fields='name', supportsAllDrives=True).execute()
                            nom_pdf = meta.get('name', '')
                        except Exception:
                            pass

                # Fallbacks canónicos si faltan nombres limpios
                tipo_c = str(row[4]).strip() if len(row) > 4 else "COM"
                tipo_cod_b = "COM" if "comercial" in tipo_c.lower() else "SER"
                dest_b = str(row[2]).strip() if len(row) > 2 else ""
                if not dest_b:
                    emp_raw = str(row[1]).strip() if len(row) > 1 else ""
                    dest_b = emp_raw.split()[0] if emp_raw else "GENERAL"

                if not nom_doc and corr_row:
                    nom_doc = f"CERT-{tipo_cod_b}-{corr_row}-{dest_b}.docx"

                if not nom_pdf and corr_row:
                    tag_rem_b = "-REM-TRAN" if (len(row) > 5 and str(row[5]).strip()) else ""
                    nom_pdf = f"CERT-{tipo_cod_b}-{corr_row}{tag_rem_b}-{dest_b}.pdf"

                item_res = {
                    "fila": num_fila,
                    "fecha": str(row[0]).strip() if len(row) > 0 else "",
                    "empresa": str(row[1]).strip() if len(row) > 1 else "",
                    "fundo": str(row[2]).strip() if len(row) > 2 else "",
                    "correlativo": corr_row,
                    "tipo_cert": str(row[4]).strip() if len(row) > 4 else "",
                    "guias": str(row[5]).strip() if len(row) > 5 else "",
                    "link_guia": raw_guia,
                    "link_doc": url_doc,
                    "raw_doc": nom_doc,
                    "link_pdf": url_pdf,
                    "raw_pdf": nom_pdf,
                    "observacion": str(row[9]).strip() if len(row) > 9 else ""
                }
                if es_exacto:
                    resultados_exactos.append(item_res)
                else:
                    resultados_parciales.append(item_res)

        return resultados_exactos + resultados_parciales
    except Exception as e:
        logger.error(f"Error buscando certificado en 'historial': {e}")
        return []

def actualizar_bitacora_guias_recibidas(guias_lista: List[str], filas_repositorio: Optional[List[int]] = None, sheets_service=None) -> int:
    """
    Registra la marca '✅ Nuevo: dd/mm/yyyy hh:mm' en la Columna H de 'Guias_recibidas'
    para todas las guías procesadas, ya sea por número de fila o buscando por número de guía.
    """
    if not sheets_service:
        _, sheets_service = obtener_servicios_google()

    ahora_pe = (datetime.now(PET)).strftime("%d/%m/%Y %H:%M")
    marca = f"✅ Nuevo: {ahora_pe}"
    filas_a_actualizar = set()

    if filas_repositorio:
        for f in filas_repositorio:
            try:
                f_int = int(f)
                if f_int >= 2:
                    filas_a_actualizar.add(f_int)
            except Exception:
                pass

    if guias_lista:
        try:
            from bot.handlers import match_guia_en_texto
            res = sheets_service.spreadsheets().values().get(
                spreadsheetId=ID_SHEET_REPOSITORIO, range="'Guias_recibidas'!B2:H"
            ).execute()
            filas_sheet = res.get('values', [])
            for idx, r in enumerate(filas_sheet):
                f_num = idx + 2
                g_sheet = r[0] if len(r) > 0 else ""
                col_h = r[6] if len(r) > 6 else ""
                if "✅" in col_h:
                    continue
                for g_buscada in guias_lista:
                    if g_sheet and match_guia_en_texto(g_buscada, g_sheet):
                        filas_a_actualizar.add(f_num)
                        break
        except Exception as e:
            logger.error(f"Error escaneando Guias_recibidas para marcar bitácora: {e}")

    actualizados = 0
    for f_num in sorted(list(filas_a_actualizar)):
        try:
            sheets_service.spreadsheets().values().update(
                spreadsheetId=ID_SHEET_REPOSITORIO,
                range=f"'Guias_recibidas'!H{f_num}",
                valueInputOption="USER_ENTERED",
                body={"values": [[marca]]}
            ).execute()
            actualizados += 1
        except Exception as e:
            logger.error(f"Error actualizando Col H en Guias_recibidas fila {f_num}: {e}")

    logger.info(f"Bitácora Guias_recibidas actualizada: {actualizados} filas marcadas con '{marca}'.")
    return actualizados

# ====================================================================
# --- BLOQUE 10: GENERACIÓN COMPLETA DE CERTIFICADOS ---
# ====================================================================
def procesar_generacion_certificado(payload: Dict[str, Any], remision_bytes: Optional[bytes] = None, transporte_bytes: Optional[bytes] = None) -> Dict[str, Any]:
    """
    Orquesta la emisión de un certificado completo:
    1. Resuelve plantilla de Drive
    2. Renderiza DocxTemplate con datos de contexto
    3. Inyecta tabla [[TABLA_NOTAS]]
    4. Convierte a PDF vía Drive API
    5. Consolida PDF unificado (Certificado + Guías)
    6. Sube Word y PDF a Drive
    7. Registra fila en 'historial' de Google Sheets
    8. Actualiza bitácora en Guias_recibidas (Columna H)
    """
    drive_service, sheets_service = obtener_servicios_google()

    empresa_firma = payload.get("empresa_firma", "EPMI S.A.C.")
    raw_flujo = str(payload.get("tipo_flujo", "")).strip()
    raw_tipo_cert = str(payload.get("tipo_certificado", "")).strip()
    flujo_detectado = raw_flujo or raw_tipo_cert or "Comercialización"
    norm_flujo = normalizar_texto_sin_tildes(flujo_detectado).lower()

    if "comercial" in norm_flujo:
        tipo_flujo = "Comercialización"
    elif "final 2" in norm_flujo or "2" in norm_flujo:
        tipo_flujo = "Disposición Final 2"
    elif "disposic" in norm_flujo or "final" in norm_flujo or "servicio" in norm_flujo:
        tipo_flujo = "Disposición Final 1"
    else:
        tipo_flujo = "Comercialización"

    tipo_cod = "COM" if "comercializa" in str(tipo_flujo).lower() else "SER"
    sec_key = "COMERCIALIZACION" if tipo_cod == "COM" else "SERVICIOS"

    # Obtener catálogo para frases oficiales por defecto
    try:
        cat_servicios = obtener_catalogo_servicios(sheets_service)
    except Exception:
        cat_servicios = {}

    def_titulos = cat_servicios.get(sec_key, {}).get("titulos", [])
    def_servicios = cat_servicios.get(sec_key, {}).get("servicios", [])
    def_residuos = cat_servicios.get(sec_key, {}).get("residuos", [])

    fallback_tit = def_titulos[0] if def_titulos else (
        "CERTIFICADO DE RECOLECCIÓN, TRANSPORTE Y \nCOMERCIALIZACIÓN DE RESIDUOS SÓLIDOS" if tipo_cod == "COM" 
        else "CERTIFICADO DE RECOLECCIÓN Y TRANSPORTE \nDE RESIDUOS SÓLIDOS PELIGROSOS"
    )
    fallback_serv = def_servicios[0] if def_servicios else (
        "RECOLECCIÓN, TRANSPORTE Y COMERCIALIZACIÓN DE RESIDUOS SÓLIDOS" if tipo_cod == "COM"
        else "RECOLECCIÓN Y TRANSPORTE DE RESIDUOS SÓLIDOS"
    )
    fallback_res = def_residuos[0] if def_residuos else (
        "RESIDUOS APROVECHABLES" if tipo_cod == "COM" else "RESIDUOS NO APROVECHABLES"
    )

    es_modelo = bool(payload.get("es_modelo", False))

    v_corr = str(payload.get("correlativo", "")).strip()

    v_tit_raw = str(payload.get("titulo", "")).strip()
    if not v_tit_raw or v_tit_raw.upper() in ["S/D", "NONE", "CERTIFICADO DE OPERACIÓN", "CERTIFICADO DE OPERACION"]:
        v_tit = fallback_tit
    else:
        v_tit = v_tit_raw
    v_tit = formatear_titulo_oficial(v_tit, tipo_cod=tipo_cod)

    v_cli = str(payload.get("cliente", "")).strip()
    v_ruc_c = str(payload.get("ruc_cliente", "")).strip()

    v_serv_raw = str(payload.get("servicio", "")).strip()
    if not v_serv_raw or v_serv_raw.upper() in ["S/D", "NONE", "COMERCIALIZACIÓN", "COMERCIALIZACION", "DISPOSICIÓN FINAL", "DISPOSICION FINAL"]:
        v_serv = fallback_serv
    else:
        v_serv = v_serv_raw

    v_res_raw = str(payload.get("tipo_residuo", "")).strip()
    if not v_res_raw or v_res_raw.upper() in ["S/D", "NONE"]:
        v_res = fallback_res
    else:
        v_res = v_res_raw

    v_partida = str(payload.get("punto_partida", "")).strip()
    # En el sistema de certificados actualizado, la dirección de llegada se envía vacía para usar la predeterminada
    v_llegada = ""

    # Regla de fecha de emisión (Prosembra / Villacurí en Comercialización => Hoy)
    v_cli_norm = normalizar_texto_sin_tildes(v_cli).upper()
    es_prosembra_villacuri = ("PROSEMBRA" in v_cli_norm) or ("VILLACURI" in v_cli_norm.replace(" ", ""))
    if es_prosembra_villacuri and tipo_cod == "COM":
        v_fec_emis = (datetime.now(PET)).strftime("%d/%m/%Y")
    else:
        v_fec_emis = str(payload.get("fecha_emision", datetime.now(PET).strftime("%d/%m/%Y"))).strip()
    items = payload.get("items", [])

    # Obtener RUC y registro de la empresa firmante
    cat_empresas = obtener_catalogo_empresas(sheets_service)
    emisor_info = cat_empresas.get(empresa_firma.upper(), {})
    emisor_ruc = emisor_info.get("ruc", "20601369792")
    emisor_reg = emisor_info.get("reg", "EO-RS-00043-2021-MINAM")
    emisor_nombre = empresa_firma

    # 1. Descargar plantilla
    plantilla_id = resolver_plantilla_drive(empresa_firma, tipo_flujo, drive_service, es_modelo)
    plantilla_bytes = descargar_plantilla_docx(plantilla_id, drive_service)

    # 2. Renderizar DocxTemplate
    doc = DocxTemplate(io.BytesIO(plantilla_bytes))
    ctx = {
        "CORRELATIVO": v_corr,
        "TITULO": v_tit,
        "REGISTRO": emisor_reg,
        "CLIENTE": v_cli,
        "RUC_CLIENTE": v_ruc_c,
        "RAZON_SOCIAL_CLIENTE": v_cli,
        "SERVICIO_O_COMPRA": v_serv,
        "SERVICIO O COMPRA": v_serv,
        "TIPO_DE_RESIDUO": v_res,
        "PUNTO_PARTIDA": v_partida,
        "DIRECCION_EMPRESA": v_llegada,
        "DIRECCION_LLEGADA": v_llegada,
        "LLEGADA": v_llegada,
        "EMPRESA_2": emisor_nombre,
        "FECHA_EMISION": v_fec_emis,
        "DESTINATARIO_FINAL": emisor_nombre,
        "EMPRESA": emisor_nombre,
        "RUC_EMPRESA": emisor_ruc,
        "RUC": emisor_ruc,
        "EMISOR": emisor_nombre,
        "RUC_EMISOR": emisor_ruc
    }
    doc.render(ctx)
    buf_tpl = io.BytesIO()
    doc.save(buf_tpl)
    tpl_bytes = buf_tpl.getvalue()
    buf_tpl.close()
    del doc

    # 3. Inyectar tabla de ítems con columnas personalizadas
    columnas_payload = payload.get("columnas")
    tpl_io = io.BytesIO(tpl_bytes)
    docx_final_bytes = inyectar_tabla_en_docx(tpl_io, items, columnas=columnas_payload)
    tpl_io.close()
    del tpl_bytes

    # 4. Convertir Certificado a PDF
    cert_pdf_bytes = convertir_docx_a_pdf(docx_final_bytes, drive_service)

    # 5. Unificar con guías si existen
    tiene_guias = bool(remision_bytes or transporte_bytes)
    if tiene_guias:
        pdf_unificado_bytes = unir_tres_documentos_pdf(cert_pdf_bytes, remision_bytes, transporte_bytes)
    else:
        pdf_unificado_bytes = cert_pdf_bytes

    # 6. Nombres de archivo y destino
    tipo_cod = "COM" if "comercializa" in str(tipo_flujo).lower() else "SER"
    destino_raw = str(v_partida).split(' - ')[-1].strip()
    destino_raw = re.sub(r'(?i)^(Av\.|Avenida|Calle|Jr\.|Jirón|Pasaje|Carretera|Panamericana)\s+', '', destino_raw).strip()
    destino_limpio = re.sub(r'(?i)^(Planta|Fundo|Sede|Sucursal|Predio)\s+', '', destino_raw).strip().upper()
    if not destino_limpio or destino_limpio == "NAN":
        destino_final = str(v_cli).strip().upper()
    else:
        destino_final = destino_limpio

    if ' - ' not in str(v_partida) and len(destino_final.split()) > 1:
        destino_final = destino_final.split()[0]

    nombre_base = f"CERT-{tipo_cod}-{v_corr}-{destino_final}"
    nombre_docx = f"{nombre_base}.docx"
    if tiene_guias:
        nombre_pdf = f"CERT-{tipo_cod}-{v_corr}-REM-TRAN-{destino_final}.pdf"
    else:
        nombre_pdf = f"{nombre_base}.pdf"

    # 7. Subir Word a Drive
    file_id_doc, link_doc = subir_a_drive_certificado(
        docx_final_bytes, nombre_docx, tipo_flujo, empresa_firma, es_modelo,
        mime_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    )

    # 8. Subir PDF a Drive
    file_id_pdf, link_pdf = subir_a_drive_certificado(
        pdf_unificado_bytes, nombre_pdf, tipo_flujo, empresa_firma, es_modelo,
        mime_type='application/pdf'
    )

    # 9. Extraer guías consolidadas para la bitácora
    guias_lista = [limpiar_numero_guia(it.get('guia_origen', '')) for it in items if limpiar_numero_guia(it.get('guia_origen', ''))]
    val_guia_completa = ", ".join(sorted(list(set(guias_lista)))) if guias_lista else limpiar_numero_guia(payload.get('guia', ''))

    val_cert = "COMERCIALIZACIÓN" if "comercializa" in str(tipo_flujo).lower() else "FINAL"
    if es_modelo:
        val_cert = f"M-{val_cert[:3]}"

    # Regla: cuando solo se tiene certificado va en columna H, y cuando se juntan certificado y guias va el PDF en columna I
    val_col_h = f'=HYPERLINK("{link_doc}"; "{nombre_docx}")' if link_doc else nombre_docx
    if tiene_guias:
        val_col_i = f'=HYPERLINK("{link_pdf}"; "{nombre_pdf}")' if link_pdf else nombre_pdf
    else:
        val_col_i = ""

    u_raw = payload.get('usuario_email') or payload.get('usuario') or 'Usuario'
    u_clean = str(u_raw).strip()
    if u_clean.startswith('@'):
        u_clean = u_clean[1:]

    fila_existente = buscar_fila_certificado_existente(v_corr, tipo_cod=tipo_cod, sheets_service=sheets_service)
    ahora_pe = datetime.now(PET).strftime("%d/%m/%Y %H:%M")

    if fila_existente:
        obs_col_j = f"Modificado por Mini App ({ahora_pe} - {u_clean})"
    else:
        obs_col_j = f"Emitido por Mini App ({ahora_pe} - {u_clean})"

    fecha_registro = datetime.now(PET).strftime("%d/%m/%Y")
    datos_log = [
        fecha_registro,
        str(v_cli).strip().upper(),
        str(destino_final).strip().upper(),
        v_corr,
        val_cert,
        val_guia_completa,
        "",  # Link guía
        val_col_h,
        val_col_i,
        obs_col_j
    ]

    if fila_existente:
        actualizar_fila_historial_sheets(fila_existente, datos_log, sheets_service=sheets_service)
        fila_historial = fila_existente
        logger.info(f"Fila {fila_existente} de 'historial' sobrescrita para certificado {tipo_cod}-{v_corr}")
    else:
        fila_historial = registrar_en_historial_sheets(datos_log, sheets_service)
        logger.info(f"Nueva fila {fila_historial} insertada en 'historial' para certificado {tipo_cod}-{v_corr}")

    # 10. Actualizar bitácora en Guias_recibidas (Columna H)
    filas_repo = payload.get('filas_repositorio') or []
    try:
        actualizar_bitacora_guias_recibidas(guias_lista, filas_repositorio=filas_repo, sheets_service=sheets_service)
    except Exception as e_bit:
        logger.error(f"Error actualizando bitácora en Guias_recibidas: {e_bit}")

    return {
        "status": "success",
        "correlativo": v_corr,
        "nombre_archivo": nombre_pdf,
        "doc_link": link_doc,
        "pdf_link": link_pdf,
        "file_id_pdf": file_id_pdf,
        "file_id_doc": file_id_doc,
        "fila_historial": fila_historial,
        "pdf_bytes": pdf_unificado_bytes,
        "cliente": v_cli,
        "fecha": v_fec_emis,
        "tipo_cod": tipo_cod,
        "guias": guias_lista,
        "filas_repositorio": filas_repo
    }

# ====================================================================
# --- BLOQUE 11: REGENERACIÓN QUIRÚRGICA DEL EXPEDIENTE ---
# ====================================================================
def procesar_regeneracion_expediente(
    correlativo: str,
    nuevo_doc_bytes: Optional[bytes] = None,
    nuevo_pdf_bytes: Optional[bytes] = None,
    nombre_pdf_usuario: str = "",
    obs_extra: str = "",
    usuario_editor: str = "Usuario MiniApp"
) -> Dict[str, Any]:
    """
    Ejecuta la sustitución quirúrgica de la carátula en el PDF consolidado existente en Drive:
    - Conserva intactas las guías escaneadas anexas.
    - Sobrescribe in-place en Drive conservando el enlace webViewLink público.
    - Actualiza columnas I y J de la pestaña 'historial'.
    """
    drive_service, sheets_service = obtener_servicios_google()

    # 1. Localizar el certificado en Historial
    certificados = buscar_datos_certificado_en_historial(correlativo, sheets_service, drive_service)
    if not certificados:
        raise ValueError(f"No se encontró ningún certificado con el correlativo '{correlativo}' en Historial.")

    cert_sel = certificados[0]
    num_fila = cert_sel['fila']
    link_doc = cert_sel.get('link_doc') or cert_sel.get('raw_doc')
    link_pdf = cert_sel.get('link_pdf') or cert_sel.get('raw_pdf')

    if not link_pdf:
        raise ValueError(f"El certificado #{correlativo} no tiene un enlace a PDF consolidado en Historial.")

    # 2. Obtener bytes del nuevo certificado (Página 1)
    if nuevo_pdf_bytes:
        nuevo_cert_pdf_bytes = nuevo_pdf_bytes
    elif nuevo_doc_bytes:
        nuevo_cert_pdf_bytes = convertir_docx_a_pdf(nuevo_doc_bytes, drive_service)
    else:
        # Descargar Word actualizado directamente desde Google Drive
        doc_id = extraer_id_drive(link_doc)
        if not doc_id and cert_sel.get('raw_doc'):
            try:
                res_f = drive_service.files().list(
                    q=f"name contains '{cert_sel['raw_doc']}' and trashed = false",
                    spaces='drive', orderBy='modifiedTime desc', fields='files(id)'
                ).execute()
                files_found = res_f.get('files', [])
                if files_found:
                    doc_id = files_found[0]['id']
            except Exception:
                pass

        if not doc_id:
            raise ValueError(f"No se pudo resolver el ID del documento Word para #{correlativo}.")
        doc_bytes = descargar_plantilla_docx(doc_id, drive_service)
        nuevo_cert_pdf_bytes = convertir_docx_a_pdf(doc_bytes, drive_service)

    # 3. Descargar PDF unificado existente de Drive asegurando la última versión
    pdf_id = extraer_id_drive(link_pdf)
    if not pdf_id and cert_sel.get('raw_pdf'):
        try:
            res_pdf = drive_service.files().list(
                q=f"name contains '{cert_sel['raw_pdf']}' and trashed = false",
                spaces='drive', orderBy='modifiedTime desc', fields='files(id)'
            ).execute()
            files_pdf = res_pdf.get('files', [])
            if files_pdf:
                pdf_id = files_pdf[0]['id']
        except Exception:
            pass

    if not pdf_id:
        raise ValueError(f"No se pudo resolver el ID de Drive del PDF consolidado para #{correlativo}.")

    try:
        rev_res = drive_service.revisions().list(
            fileId=pdf_id, fields='revisions(id, modifiedTime)'
        ).execute()
        revs = rev_res.get('revisions', [])
        if revs:
            latest_rev = sorted(revs, key=lambda r: r.get('modifiedTime', ''))[-1]
            logger.info(f"Descargando última revisión {latest_rev.get('id')} ({latest_rev.get('modifiedTime')}) del PDF consolidado {pdf_id}")
            req = drive_service.revisions().get_media(
                fileId=pdf_id, revisionId=latest_rev['id']
            )
        else:
            req = drive_service.files().get_media(fileId=pdf_id, supportsAllDrives=True)
    except Exception as e_rev:
        logger.warning(f"Error consultando revisiones del PDF consolidado {pdf_id} ({e_rev}), usando get_media estándar.")
        req = drive_service.files().get_media(fileId=pdf_id, supportsAllDrives=True)

    fh = io.BytesIO()
    dl = MediaIoBaseDownload(fh, req)
    done = False
    while not done:
        _, done = dl.next_chunk()
    pdf_unificado_existente_bytes = fh.getvalue()
    fh.close()

    # 4. Sustitución quirúrgica de la carátula
    pdf_actualizado_bytes = sustituir_certificado_en_pdf(
        pdf_unificado_existente_bytes, nuevo_cert_pdf_bytes, num_paginas_reemplazar=1
    )

    # 5. Nombre final canónico
    tipo_cod = "COM" if "comercial" in str(cert_sel.get('tipo_cert', '')).lower() else "SER"
    destino = str(cert_sel.get('fundo') or cert_sel.get('empresa') or '').strip().upper()
    destino_raw = destino.split(' - ')[-1].strip()
    destino_raw = re.sub(r'(?i)^(Av\.|Avenida|Calle|Jr\.|Jirón|Pasaje|Carretera|Panamericana)\s+', '', destino_raw).strip()
    destino_limpio = re.sub(r'(?i)^(Planta|Fundo|Sede|Sucursal|Predio)\s+', '', destino_raw).strip().upper()
    if not destino_limpio:
        destino_limpio = str(cert_sel.get('empresa', '')).split()[0].upper() if cert_sel.get('empresa') else "GENERAL"

    es_unificado = "REM-TRAN" in str(cert_sel.get('raw_pdf', '')).upper() or bool(cert_sel.get('guias'))
    tag_rem = "-REM-TRAN" if es_unificado else ""

    if nombre_pdf_usuario and not nombre_pdf_usuario.strip().startswith(('http://', 'https://')):
        nombre_final = nombre_pdf_usuario.strip()
    else:
        nombre_final = f"CERT-{tipo_cod}-{correlativo}{tag_rem}-{destino_limpio}.pdf"

    if not nombre_final.lower().endswith('.pdf'):
        nombre_final += '.pdf'

    # 6. Sobrescribir in-place en Drive
    nuevo_link_drive = sobrescribir_o_subir_pdf_drive(
        pdf_id, pdf_actualizado_bytes, nombre_final,
        cert_sel.get('tipo_cert', 'Comercialización'), cert_sel.get('empresa', 'EPMI S.A.C.')
    )

    # 7. Registrar auditoría en Sheets con fórmula HYPERLINK
    link_para_historial = nuevo_link_drive or link_pdf
    val_col_i = f'=HYPERLINK("{link_para_historial}"; "{nombre_final}")'
    registrar_edicion_en_historial(
        num_fila, val_col_i, usuario_editor=usuario_editor,
        obs_extra=obs_extra, sheets_service=sheets_service
    )
    liberar_memoria_proceso()

    return {
        "status": "success",
        "correlativo": correlativo,
        "fila": num_fila,
        "pdf_link": link_para_historial,
        "nombre_archivo": nombre_final,
        "pdf_bytes": pdf_actualizado_bytes
    }

# ====================================================================
# --- BLOQUE 12: CONSULTA DE GUÍAS PENDIENTES (REPOSITORIO) ---
# ====================================================================
def obtener_guias_pendientes_repositorio(sheets_service=None) -> List[Dict[str, Any]]:
    """
    Retorna la lista de guías registradas en 'Guias_recibidas' que aún no tienen certificado emitido.
    Cruza contra 'historial' y contra columna H para evitar duplicados.
    Resuelve RUC y Dirección desde los catálogos de CLIENTES y Direcciones.
    """
    if not sheets_service:
        _, sheets_service = obtener_servicios_google()
    try:
        # 1. Cargar guías emitidas en 'historial'
        res_hist = sheets_service.spreadsheets().values().get(
            spreadsheetId=ID_SHEET_CONTROL, range="'historial'!F2:F"
        ).execute()
        filas_hist = res_hist.get('values', [])
        guias_emitidas_hist = [str(r[0]).strip().upper() for r in filas_hist if r and len(r) > 0 and str(r[0]).strip()]

        # 2. Cargar catálogos para enriquecer cada guía
        cat_clientes = obtener_catalogo_clientes_certificados(sheets_service)
        cat_direcciones = obtener_catalogo_direcciones_certificados(sheets_service)

        # 3. Importar normalizador / matcher
        from bot.handlers import match_guia_en_texto

        # 4. Leer 'Guias_recibidas'
        res_recib = sheets_service.spreadsheets().values().get(
            spreadsheetId=ID_SHEET_REPOSITORIO, range="'Guias_recibidas'!A2:K"
        ).execute()
        filas_recib = res_recib.get('values', [])

        pendientes = []
        for i, fila in enumerate(filas_recib):
            fila_num = i + 2
            if len(fila) >= 4:
                fecha_str = str(fila[0]).strip() if len(fila) > 0 else ""
                guia_num = str(fila[1]).strip() if len(fila) > 1 else ""
                tipo_guia = str(fila[2]).strip() if len(fila) > 2 else ""
                empresa = str(fila[3]).strip() if len(fila) > 3 else ""
                fundo = str(fila[4]).strip() if len(fila) > 4 else ""
                archivo = str(fila[5]).strip() if len(fila) > 5 else ""
                col_cert = str(fila[7]).strip() if len(fila) > 7 else ""
                tipo_cert = str(fila[9]).strip() if len(fila) > 9 else ""

                if not guia_num or not empresa or empresa.upper() in ["S/D", "EMPRESA"]:
                    continue

                # Si ya tiene un certificado emitido o marca en Col H (ej. '✅ Nuevo:...'):
                if col_cert and ("✅" in col_cert or "NUEVO" in col_cert.upper() or any(term in col_cert.upper() for term in ["EMITIDO", "CERT", "FINAL", "COMERCIAL"]) or len(col_cert.strip()) > 3):
                    continue

                # Si ya está registrado en la pestaña historial:
                if any(match_guia_en_texto(guia_num, gh) for gh in guias_emitidas_hist):
                    continue

                # Resolver RUC
                empresa_upper = empresa.upper()
                ruc_resuelto = ""
                if empresa_upper in cat_clientes:
                    ruc_resuelto = cat_clientes[empresa_upper].get("ruc", "")
                else:
                    for k_cli, v_cli in cat_clientes.items():
                        if k_cli in empresa_upper or empresa_upper in k_cli:
                            ruc_resuelto = v_cli.get("ruc", "")
                            break

                # Resolver Dirección de Partida
                fundo_upper = fundo.upper()
                dir_resuelta = ""
                dirs_empresa = cat_direcciones.get(empresa_upper, {})
                if not dirs_empresa:
                    for k_dir, v_map in cat_direcciones.items():
                        if k_dir in empresa_upper or empresa_upper in k_dir:
                            dirs_empresa = v_map
                            break
                if dirs_empresa:
                    if fundo_upper in dirs_empresa:
                        dir_resuelta = dirs_empresa[fundo_upper]
                    else:
                        for f_k, d_val in dirs_empresa.items():
                            if f_k in fundo_upper or fundo_upper in f_k:
                                dir_resuelta = d_val
                                break

                if not dir_resuelta and fundo:
                    dir_resuelta = f"Fundo - {fundo}"

                # Normalizar tipo sugerido exactamente a las opciones del selector
                tipo_raw_norm = normalizar_texto_sin_tildes(tipo_cert).upper()
                if "COMERC" in tipo_raw_norm:
                    tipo_sug = "Comercialización"
                elif "FINAL 2" in tipo_raw_norm or "2" in tipo_raw_norm:
                    tipo_sug = "Disposición Final 2"
                elif "DISPOSIC" in tipo_raw_norm or "FINAL" in tipo_raw_norm or "SERVICIO" in tipo_raw_norm:
                    tipo_sug = "Disposición Final 1"
                else:
                    tipo_sug = "Disposición Final 1" if "PETRAMAS" in empresa_upper else "Comercialización"

                pendientes.append({
                    "fila": fila_num,
                    "fecha": fecha_str,
                    "empresa": empresa,
                    "fundo": fundo,
                    "guia": limpiar_numero_guia(guia_num),
                    "tipo_guia": tipo_guia,
                    "ruc": ruc_resuelto,
                    "direccion_partida": dir_resuelta,
                    "archivo": archivo,
                    "tipo_sugerido": tipo_sug
                })

        return pendientes
    except Exception as e:
        logger.error(f"Error consultando guías pendientes: {e}")
        return []

def consolidar_guias_repositorio_ocr(guias: List[Dict[str, Any]], drive_service=None) -> Dict[str, Any]:
    """
    Descarga cada guía de Drive seleccionada, ejecuta Vertex OCR,
    extrae vehículos (placas), ítems con sus descripciones/cantidades/pesos,
    y unifica todo para el formulario de la Mini App.
    """
    if not drive_service:
        drive_service, _ = obtener_servicios_google()

    items_resultado = []
    placas = []
    punto_partida_detectado = ""
    punto_llegada_detectado = ""
    cliente_detectado = ""
    ruc_detectado = ""

    for g in guias:
        nombre_o_id = g.get('archivo') or g.get('link_guia') or ''
        guia_num_fallback = limpiar_numero_guia(g.get('guia') or g.get('numero_guia') or '')
        fecha_fallback = normalizar_fecha(g.get('fecha', ''))

        pdf_bytes = None
        if nombre_o_id:
            pdf_bytes = descargar_archivo_drive_por_id_o_nombre(nombre_o_id, drive_service)

        datos_ocr = None
        if pdf_bytes:
            try:
                datos_ocr = procesar_guia_ia_vertex(pdf_bytes)
            except Exception as e:
                logger.warning(f"Error OCR para guía '{guia_num_fallback}': {e}")

        if datos_ocr:
            s = limpiar_numero_guia(datos_ocr.get('serie') or guia_num_fallback)
            f = normalizar_fecha(datos_ocr.get('fecha') or fecha_fallback)
            p = str(datos_ocr.get('vehiculo', '')).strip().upper()

            # Auto-corrección si Vertex invirtió serie y placa
            if es_formato_placa(s) and es_formato_guia(p):
                s, p = p, s
                s = limpiar_numero_guia(s)

            if p and p not in placas:
                placas.append(p)

            if not punto_partida_detectado and datos_ocr.get('punto_partida'):
                punto_partida_detectado = datos_ocr.get('punto_partida')

            if not punto_llegada_detectado and datos_ocr.get('punto_llegada'):
                punto_llegada_detectado = datos_ocr.get('punto_llegada')

            if not cliente_detectado and datos_ocr.get('cliente'):
                cliente_detectado = datos_ocr.get('cliente')

            if not ruc_detectado and datos_ocr.get('ruc_cliente'):
                ruc_detectado = datos_ocr.get('ruc_cliente')

            ocr_items = datos_ocr.get('items', [])
            if ocr_items:
                for it in ocr_items:
                    desc_limpia = limpiar_descripcion(it.get('desc', ''))
                    cant_limpia = formato_inteligente(limpiar_monto(it.get('cant', 1)))
                    um_raw = str(it.get('um', 'KG')).upper()
                    um_limpia = 'KG' if 'KILO' in um_raw else ('GLN' if 'GALO' in um_raw else ('UNID' if 'UNIDA' in um_raw else um_raw))
                    peso_limpio = formato_inteligente(limpiar_monto(it.get('peso', 0)))

                    items_resultado.append({
                        'fecha_origen': f,
                        'placa_origen': p,
                        'guia_origen': s,
                        'desc': desc_limpia,
                        'cant': cant_limpia,
                        'um': um_limpia,
                        'peso': peso_limpio
                    })
            else:
                items_resultado.append({
                    'fecha_origen': f,
                    'placa_origen': p,
                    'guia_origen': s,
                    'desc': 'RESIDUOS RECICLABLES Y APROVECHABLES',
                    'cant': '1',
                    'um': 'KG',
                    'peso': '0.00'
                })
        else:
            # Fallback si no hay PDF o falló OCR
            items_resultado.append({
                'fecha_origen': fecha_fallback,
                'placa_origen': '',
                'guia_origen': guia_num_fallback,
                'desc': 'RESIDUOS RECICLABLES Y APROVECHABLES',
                'cant': '1',
                'um': 'KG',
                'peso': '0.00'
            })

    def fecha_a_entero(fecha_str):
        try:
            p = str(fecha_str).strip().split('/')
            if len(p) == 3:
                return int(f"{p[2]}{p[1]}{p[0]}")
        except Exception:
            pass
        return 99999999

    items_resultado.sort(key=lambda x: fecha_a_entero(x['fecha_origen']))
    placa_sugerida = ", ".join(placas) if placas else ""

    return {
        "success": True,
        "placa": placa_sugerida,
        "placas": placas,
        "items": items_resultado,
        "punto_partida": punto_partida_detectado,
        "punto_llegada": punto_llegada_detectado,
        "cliente": cliente_detectado,
        "ruc_cliente": ruc_detectado,
        "total_procesadas": len(guias)
    }

def obtener_url_webapp_certificados(correlativo: Optional[str] = None, datos_edicion: Optional[Dict[str, Any]] = None, user_id: Optional[Any] = None, modo: Optional[str] = None) -> str:
    """Construye la URL segura HTTPS para abrir la Telegram Mini App de Certificados."""
    import urllib.parse
    import base64
    from config.settings import WEBAPP_CERTIFICADOS_URL, RENDER_EXTERNAL_URL

    base_url = WEBAPP_CERTIFICADOS_URL
    if not base_url and RENDER_EXTERNAL_URL:
        base_url = f"{RENDER_EXTERNAL_URL.rstrip('/')}/certificados"
    if not base_url:
        base_url = "https://lia-w2yb.onrender.com/certificados"

    params = {'v': str(int(time.time()))}
    if correlativo:
        params['corr'] = str(correlativo)
    if user_id:
        params['uid'] = str(user_id)
    if modo:
        params['modo'] = str(modo)
    if datos_edicion:
        json_str = json.dumps(datos_edicion, ensure_ascii=False)
        params['data'] = base64.b64encode(json_str.encode('utf-8')).decode('utf-8')

    sep = '&' if '?' in base_url else '?'
    return f"{base_url}{sep}{urllib.parse.urlencode(params)}"

# ====================================================================
# --- BLOQUE 14: JUNTE DE EXPEDIENTE (CERTIFICADO + GUÍAS) ---
# ====================================================================
def buscar_guias_asociadas_para_unir(guias_lista: List[str], sheets_service=None, drive_service=None) -> Tuple[List[str], List[str]]:
    """
    Localiza los archivos de Google Drive (nombres o IDs) correspondientes a la(s) Guía(s) de Remisión
    y a la(s) Guía(s) de Transporte cruzando 'Registro_Guias', 'Guias_recibidas' y búsqueda directa en Drive.
    Retorna (archivos_remision, archivos_transporte).
    """
    if not drive_service or not sheets_service:
        drive_service, sheets_service = obtener_servicios_google()

    from bot.handlers import match_guia_en_texto

    archivos_remision = []
    archivos_transporte = []
    guias_ligadas_descubiertas = []

    # 1. Escanear Registro_Guias (Columnas A:N)
    try:
        res_reg = sheets_service.spreadsheets().values().get(
            spreadsheetId=ID_SHEET_CONTROL, range="Registro_Guias!A2:N"
        ).execute()
        filas_reg = res_reg.get('values', [])
    except Exception as e:
        logger.error(f"Error leyendo Registro_Guias: {e}")
        filas_reg = []

    for r in filas_reg:
        num_guia = r[1] if len(r) > 1 else ''
        guia_lig = r[2] if len(r) > 2 else ''
        tipo_guia = (r[3] if len(r) > 3 else '').upper()
        archivo_hecho = r[8] if len(r) > 8 else ''
        archivo_recib = r[9] if len(r) > 9 else ''
        archivo = archivo_hecho or archivo_recib
        obs = (r[11] if len(r) > 11 else '').upper()

        if 'ERRADA' in obs or 'ERROR' in obs:
            continue

        for g in guias_lista:
            if match_guia_en_texto(g, num_guia) or match_guia_en_texto(g, guia_lig):
                if 'TRANS' in tipo_guia:
                    if archivo and archivo not in archivos_transporte:
                        archivos_transporte.append(archivo)
                else:
                    if archivo and archivo not in archivos_remision:
                        archivos_remision.append(archivo)

                if guia_lig and guia_lig.upper() not in ['SIN GUIA', 'NO TIENE', '-', ''] and guia_lig not in guias_ligadas_descubiertas:
                    guias_ligadas_descubiertas.append(guia_lig)

    # 2. Para contrapartes descubiertas en Registro_Guias:
    for g_lig in guias_ligadas_descubiertas:
        for r2 in filas_reg:
            n2 = r2[1] if len(r2) > 1 else ''
            t2 = (r2[3] if len(r2) > 3 else '').upper()
            a2 = (r2[8] if len(r2) > 8 else '') or (r2[9] if len(r2) > 9 else '')
            obs2 = (r2[11] if len(r2) > 11 else '').upper()
            if 'ERRADA' in obs2 or 'ERROR' in obs2:
                continue
            if match_guia_en_texto(g_lig, n2):
                if 'TRANS' in t2 and a2 and a2 not in archivos_transporte:
                    archivos_transporte.append(a2)
                elif a2 and a2 not in archivos_remision:
                    archivos_remision.append(a2)

    # 3. Escanear Guias_recibidas (Columnas A:F)
    try:
        res_rec = sheets_service.spreadsheets().values().get(
            spreadsheetId=ID_SHEET_REPOSITORIO, range="'Guias_recibidas'!A2:F"
        ).execute()
        filas_rec = res_rec.get('values', [])
    except Exception as e:
        logger.error(f"Error leyendo Guias_recibidas para unión: {e}")
        filas_rec = []

    todas_guias_buscar = list(set(guias_lista + guias_ligadas_descubiertas))
    for r in filas_rec:
        num = r[1] if len(r) > 1 else ''
        tipo = (r[2] if len(r) > 2 else '').upper()
        arch = r[5] if len(r) > 5 else ''
        if not arch:
            continue
        for g in todas_guias_buscar:
            if match_guia_en_texto(g, num):
                if 'TRANS' in tipo:
                    if arch not in archivos_transporte:
                        archivos_transporte.append(arch)
                else:
                    if arch not in archivos_remision:
                        archivos_remision.append(arch)

    # 4. Si aún faltan archivos de transporte o remisión, buscar directamente en Drive por nombre de guía
    for g in todas_guias_buscar:
        partes = [p.strip() for p in str(g).split('-') if p.strip()]
        filtros = []
        if len(partes) >= 2:
            filtros.append(f"name contains '{partes[0]}' and name contains '{partes[1].lstrip('0')}'")
        filtros.append(f"name contains '{g}'")

        for f_q in filtros:
            try:
                q = f"({f_q}) and mimeType = 'application/pdf' and trashed = false"
                res_d = drive_service.files().list(q=q, fields='files(id, name)', supportsAllDrives=True).execute()
                for fd in res_d.get('files', []):
                    fname = fd['name']
                    if 'TR' in fname.upper() or 'TRANS' in fname.upper():
                        if fname not in archivos_transporte:
                            archivos_transporte.append(fname)
                    else:
                        if fname not in archivos_remision:
                            archivos_remision.append(fname)
            except Exception as e_dq:
                pass

    return archivos_remision, archivos_transporte

def juntar_expediente_completo(correlativo: str, guias_lista: Optional[List[str]] = None, usuario_editor: str = "MiniApp") -> Dict[str, Any]:
    """
    Une el Certificado emitido con sus Guías de Remisión y de Transporte en un solo PDF:
    Orden estricto:
      1. Certificado (Página inicial)
      2. Guía(s) de Remisión
      3. Guía(s) de Transporte
    Sobrescribe el PDF en Google Drive in-place y registra auditoría en 'historial'.
    """
    drive_service, sheets_service = obtener_servicios_google()

    # 1. Localizar el certificado en Historial
    certificados = buscar_datos_certificado_en_historial(correlativo, sheets_service, drive_service)
    if not certificados:
        raise ValueError(f"No se encontró ningún certificado con el correlativo '{correlativo}' en Historial.")

    cert_sel = certificados[0]
    num_fila = cert_sel['fila']
    link_pdf = cert_sel.get('link_pdf') or cert_sel.get('raw_pdf') or ""
    link_doc = cert_sel.get('link_doc') or cert_sel.get('raw_doc') or ""
    empresa = cert_sel.get('empresa', 'EPMI S.A.C.')
    tipo_cert = cert_sel.get('tipo_cert', 'Comercialización')
    tipo_cod = "COM" if "comercial" in str(tipo_cert).lower() else "SER"

    # Construir nombre corporativo oficial para expediente con guías (-REM-TRAN-)
    destino = str(cert_sel.get('fundo') or cert_sel.get('empresa') or '').strip().upper()
    destino_raw = destino.split(' - ')[-1].strip()
    destino_raw = re.sub(r'(?i)^(Av\.|Avenida|Calle|Jr\.|Jirón|Pasaje|Carretera|Panamericana)\s+', '', destino_raw).strip()
    destino_limpio = re.sub(r'(?i)^(Planta|Fundo|Sede|Sucursal|Predio)\s+', '', destino_raw).strip().upper()
    if not destino_limpio:
        destino_limpio = str(cert_sel.get('empresa', '')).split()[0].upper() if cert_sel.get('empresa') else "GENERAL"

    nombre_pdf = f"CERT-{tipo_cod}-{correlativo}-REM-TRAN-{destino_limpio}.pdf"

    pdf_id = extraer_id_drive(link_pdf) if link_pdf else None
    cert_pdf_bytes = None
    es_desde_docx = False

    # 2. Descargar y preparar carátula del certificado:
    # Prioridad 1: Si existe link_doc (Word en Google Drive), descargar el Word editado y convertirlo a PDF.
    # Esto asegura que cualquier modificación realizada al certificado en Drive se refleje fielmente en el expediente final.
    if link_doc:
        doc_id = extraer_id_drive(link_doc)
        if doc_id:
            try:
                doc_bytes = descargar_plantilla_docx(doc_id, drive_service)
                if doc_bytes:
                    cert_pdf_bytes = convertir_docx_a_pdf(doc_bytes, drive_service)
                    if cert_pdf_bytes:
                        es_desde_docx = True
            except Exception as e_docx:
                logger.warning(f"No se pudo convertir Word editado {doc_id} a PDF: {e_docx}")

    # Prioridad 2: Si no se pudo obtener desde Word, descargar el PDF existente de Drive
    if not cert_pdf_bytes and pdf_id:
        try:
            req = drive_service.files().get_media(fileId=pdf_id, supportsAllDrives=True)
            fh = io.BytesIO()
            dl = MediaIoBaseDownload(fh, req)
            done = False
            while not done:
                _, done = dl.next_chunk()
            cert_pdf_bytes = fh.getvalue()
        except Exception as e_dl:
            logger.warning(f"No se pudo descargar PDF previo por ID {pdf_id}: {e_dl}")

    # Prioridad 3: Buscar carátula por nombre estándar en Drive
    if not cert_pdf_bytes:
        caratula_nom = f"CERT-{tipo_cod}-{correlativo}-{destino_limpio}.pdf"
        b_caratula = descargar_archivo_drive_por_id_o_nombre(caratula_nom, drive_service)
        if b_caratula:
            cert_pdf_bytes = b_caratula

    if not cert_pdf_bytes:
        raise ValueError(f"No se pudo recuperar el PDF o Word base para el certificado #{correlativo}.")

    # 3. Determinar lista de guías a buscar
    if not guias_lista:
        raw_guias_str = cert_sel.get('guias', '')
        guias_lista = [g.strip() for g in raw_guias_str.split(',') if g.strip()]

    if not guias_lista and cert_sel.get('link_guia'):
        guias_lista = [cert_sel.get('link_guia').strip()]

    archivos_rem, archivos_trans = buscar_guias_asociadas_para_unir(guias_lista, sheets_service, drive_service)

    # 4. Descargar PDFs de guías
    bytes_remision = []
    for arch in archivos_rem:
        b = descargar_archivo_drive_por_id_o_nombre(arch, drive_service)
        if b:
            bytes_remision.append(b)

    bytes_transporte = []
    for arch in archivos_trans:
        b = descargar_archivo_drive_por_id_o_nombre(arch, drive_service)
        if b:
            bytes_transporte.append(b)

    if not bytes_remision and not bytes_transporte:
        raise ValueError(f"No se encontraron archivos PDF de guías asociados en Drive para las guías: {', '.join(guias_lista)}")

    # 5. Unir con PdfWriter:
    # Carátula del Certificado (si provino de Word editado incluye todas sus páginas; si provino de PDF previo se toma la página 1)
    # Páginas intermedias: Guías de Remisión
    # Páginas finales: Guías de Transporte
    writer = PdfWriter()
    reader_cert = PdfReader(io.BytesIO(cert_pdf_bytes))
    if es_desde_docx:
        for p in reader_cert.pages:
            writer.add_page(p)
    elif len(reader_cert.pages) > 0:
        writer.add_page(reader_cert.pages[0])

    for b_rem in bytes_remision:
        try:
            r = PdfReader(io.BytesIO(b_rem))
            for p in r.pages:
                writer.add_page(p)
        except Exception as e:
            logger.warning(f"Error añadiendo páginas de remisión: {e}")

    for b_tr in bytes_transporte:
        try:
            r = PdfReader(io.BytesIO(b_tr))
            for p in r.pages:
                writer.add_page(p)
        except Exception as e:
            logger.warning(f"Error añadiendo páginas de transporte: {e}")

    out_io = io.BytesIO()
    writer.write(out_io)
    pdf_unificado_bytes = out_io.getvalue()
    total_pags = len(writer.pages)
    out_io.close()
    del writer, reader_cert
    del bytes_remision, bytes_transporte

    # 6. Sobrescribir in-place o subir nuevo a Drive
    if pdf_id:
        nuevo_link = sobrescribir_o_subir_pdf_drive(
            pdf_id, pdf_unificado_bytes, nombre_pdf, tipo_cert, empresa
        )
    else:
        file_id_subido, nuevo_link = subir_a_drive_certificado(
            pdf_unificado_bytes, nombre_pdf, tipo_cert, empresa, mime_type='application/pdf'
        )
        pdf_id = file_id_subido

    # 7. Registrar auditoría y fórmula HYPERLINK en Sheets (Columna I y J)
    obs_unir = f"Expediente unificado: {len(archivos_rem)} Guía(s) Remisión, {len(archivos_trans)} Guía(s) Transporte"
    link_final_hist = nuevo_link or link_pdf
    val_col_i = f'=HYPERLINK("{link_final_hist}"; "{nombre_pdf}")'
    registrar_edicion_en_historial(
        num_fila, val_col_i, usuario_editor=usuario_editor,
        obs_extra=obs_unir, sheets_service=sheets_service
    )
    liberar_memoria_proceso()

    return {
        "status": "success",
        "correlativo": correlativo,
        "fila": num_fila,
        "pdf_link": link_final_hist,
        "nombre_archivo": nombre_pdf,
        "remisiones_unidas": len(archivos_rem),
        "transportes_unidos": len(archivos_trans),
        "total_paginas": total_pags,
        "pdf_bytes": pdf_unificado_bytes
    }
