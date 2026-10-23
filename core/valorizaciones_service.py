# ====================================================================
# --- MÓDULO DE VALORIZACIONES (LIQUIDACIÓN DE SERVICIOS) - EPMI SAC ---
# ====================================================================
import os
import io
import time
import re
import datetime
from datetime import timezone, timedelta
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseDownload, MediaIoBaseUpload, MediaInMemoryUpload
from google.auth.transport.requests import Request
import openpyxl

from config.settings import (
    logger,
    SHEET_VALORIZACIONES_ID,
    DRIVE_FOLDER_VALORIZACIONES_PDF,
    WEBAPP_VALORIZACIONES_URL,
    RENDER_EXTERNAL_URL
)
from core.sheets_client import obtener_credenciales

PET = timezone(timedelta(hours=-5))

def obtener_url_webapp_valorizaciones(user_id=None, modo="heredar", tab_fuente=None):
    """Construye la URL para abrir la Telegram Mini App de Valorizaciones."""
    import urllib.parse
    base_url = WEBAPP_VALORIZACIONES_URL
    if not base_url and RENDER_EXTERNAL_URL:
        base_url = f"{RENDER_EXTERNAL_URL.rstrip('/')}/valorizaciones"
    if not base_url:
        base_url = "https://lia-w2yb.onrender.com/valorizaciones"

    params = {'v': str(int(time.time()))}
    if user_id:
        params['uid'] = str(user_id)
    if modo:
        params['modo'] = str(modo)
    if tab_fuente:
        params['tab'] = str(tab_fuente)

    sep = '&' if '?' in base_url else '?'
    return f"{base_url}{sep}{urllib.parse.urlencode(params)}"

def descargar_workbook_drive(drive_service, doc_id=None, data_only=False):
    """Descarga el archivo Excel de Google Drive y lo abre con openpyxl."""
    if not doc_id:
        doc_id = SHEET_VALORIZACIONES_ID
    request = drive_service.files().get_media(fileId=doc_id)
    fh = io.BytesIO()
    downloader = MediaIoBaseDownload(fh, request)
    done = False
    while not done:
        _, done = downloader.next_chunk()
    fh.seek(0)
    wb = openpyxl.load_workbook(fh, data_only=data_only)
    return wb

def subir_workbook_drive(drive_service, wb, doc_id=None):
    """Guarda el workbook de openpyxl y actualiza el archivo en Google Drive."""
    if not doc_id:
        doc_id = SHEET_VALORIZACIONES_ID
    out_fh = io.BytesIO()
    wb.save(out_fh)
    out_fh.seek(0)
    media = MediaIoBaseUpload(
        out_fh,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        resumable=True
    )
    drive_service.files().update(fileId=doc_id, media_body=media).execute()
    logger.info(f"💾 Archivo Excel {doc_id} actualizado en Google Drive exitosamente.")

def limpiar_monto_numerico(val):
    """Convierte un valor de celda a float numérico limpio."""
    if val is None:
        return 0.0
    if isinstance(val, (int, float)):
        return float(val)
    val_str = str(val).replace("S/.", "").replace("S/", "").replace("$", "").replace(" ", "").strip()
    if "." in val_str and "," in val_str:
        val_str = val_str.replace(".", "").replace(",", ".")
    elif "," in val_str and "." not in val_str:
        val_str = val_str.replace(",", ".")
    try:
        return float(re.sub(r"[^\d.-]", "", val_str))
    except Exception:
        return 0.0

def formatear_fecha_str(val):
    """Convierte fecha de celda a formato DD/MM/YYYY."""
    if val is None:
        return "-"
    if isinstance(val, (datetime.datetime, datetime.date)):
        return f"{val.day:02d}/{val.month:02d}/{val.year}"
    val_s = str(val).strip()
    return val_s if val_s else "-"

def listar_pestanas_valorizacion():
    """
    Lista todas las pestañas disponibles en el archivo de Google Drive.
    Retorna la lista ordenada con la última pestaña creada.
    """
    try:
        creds = obtener_credenciales()
        drive = build('drive', 'v3', credentials=creds)
        wb = descargar_workbook_drive(drive, data_only=True)

        pestanas = []
        for name in wb.sheetnames:
            if name.startswith('_'):
                continue
            ws = wb[name]
            cliente = str(ws['C5'].value or "").strip()
            
            ruc_raw = ws['C6'].value
            if isinstance(ruc_raw, float):
                ruc = f"{int(ruc_raw)}"
            else:
                ruc = str(ruc_raw or "").replace(".0", "").strip()

            liq_raw = ws['C9'].value
            if isinstance(liq_raw, float):
                liq = f"{int(liq_raw):03d}"
            else:
                liq = str(liq_raw or "").replace(".0", "").strip()

            fecha = formatear_fecha_str(ws['C10'].value)

            pestanas.append({
                "id": name,
                "title": name,
                "cliente": cliente,
                "ruc": ruc,
                "liquidacion": liq,
                "fecha": fecha
            })
        return pestanas
    except Exception as e:
        logger.error(f"Error listando pestañas de valorización: {e}", exc_info=True)
        return []

def obtener_detalle_pestana(tab_name):
    """Lee y estructura los datos y filas de una pestaña de valorización."""
    try:
        creds = obtener_credenciales()
        drive = build('drive', 'v3', credentials=creds)
        wb = descargar_workbook_drive(drive, data_only=True)

        if str(tab_name) not in wb.sheetnames:
            matched = [s for s in wb.sheetnames if str(tab_name).upper() in s.upper()]
            if matched:
                tab_name = matched[0]
            else:
                tab_name = wb.sheetnames[-1]

        ws = wb[tab_name]

        cliente = str(ws['C5'].value or "").strip()
        ruc_raw = ws['C6'].value
        ruc = f"{int(ruc_raw)}" if isinstance(ruc_raw, float) else str(ruc_raw or "").replace(".0", "").strip()
        planta = str(ws['C7'].value or "").strip()
        responsable = str(ws['C8'].value or "").strip()
        liq_raw = ws['C9'].value
        liquidacion = f"{int(liq_raw):03d}" if isinstance(liq_raw, float) else str(liq_raw or "").replace(".0", "").strip()
        fecha = formatear_fecha_str(ws['C10'].value)

        datos = {
            "title": tab_name,
            "id": tab_name,
            "cliente": cliente,
            "ruc": ruc,
            "planta": planta,
            "responsable": responsable,
            "liquidacion": liquidacion,
            "fecha": fecha,
            "items": []
        }

        # Leer filas de servicios (fila 13 en adelante)
        for r in range(13, ws.max_row + 1):
            c1_val = str(ws.cell(r, 1).value or "").strip().upper()
            if "TOTAL GENERADO" in c1_val or "SUB TOTAL" in c1_val or "* METODO" in c1_val or "FORMA DE PAGO" in c1_val:
                break

            desc = str(ws.cell(r, 2).value or "").strip()
            if not desc and ws.cell(r, 1).value is None and ws.cell(r, 6).value is None:
                continue

            num_raw = ws.cell(r, 1).value
            if isinstance(num_raw, float):
                num_str = str(int(num_raw))
            else:
                num_str = str(num_raw or "").replace(".0", "").strip()

            fecha_item = formatear_fecha_str(ws.cell(r, 4).value)
            guia_item = str(ws.cell(r, 5).value or "-").strip()
            cant_item = limpiar_monto_numerico(ws.cell(r, 6).value) or 1.0
            costo_item = limpiar_monto_numerico(ws.cell(r, 7).value)
            total_item = limpiar_monto_numerico(ws.cell(r, 8).value) or (cant_item * costo_item)

            datos["items"].append({
                "num": num_str,
                "descripcion": desc,
                "fecha": fecha_item,
                "guia": guia_item,
                "cantidad": cant_item,
                "costo": costo_item,
                "total": total_item
            })

        return datos
    except Exception as e:
        logger.error(f"Error obteniendo detalle de pestaña {tab_name}: {e}", exc_info=True)
        return None

def guardar_valorizacion(datos):
    """
    Crea, hereda o modifica una pestaña en el archivo de Google Drive.
    Garantiza fórmulas automáticas (=F*G, SUM, IGV 18%, TOTAL).
    """
    try:
        creds = obtener_credenciales()
        drive = build('drive', 'v3', credentials=creds)
        wb = descargar_workbook_drive(drive, data_only=False)

        modo = datos.get("modo", "heredar")
        fuente_tab = str(datos.get("fuente_tab", "")).strip()
        nombre_pestana = str(datos.get("nombre_pestana", "")).strip()
        cliente = str(datos.get("cliente", "")).strip().upper()
        ruc = str(datos.get("ruc", "")).strip()
        planta = str(datos.get("planta", "")).strip()
        responsable = str(datos.get("responsable", "")).strip()
        liquidacion = str(datos.get("liquidacion", "")).strip()
        fecha_emision = str(datos.get("fecha", "")).strip() or datetime.datetime.now(PET).strftime("%d/%m/%Y")
        items = datos.get("items", [])

        if not nombre_pestana:
            sigla = cliente.replace("CORPORACION ", "").replace("S.A.C.", "").replace("S.A.", "").strip()[:10]
            nombre_pestana = f"{sigla}(SER {liquidacion.lstrip('0') or '1'})"

        target_ws = None

        if modo == "modificar" and nombre_pestana in wb.sheetnames:
            target_ws = wb[nombre_pestana]
        else:
            if fuente_tab in wb.sheetnames:
                source_ws = wb[fuente_tab]
            else:
                source_ws = wb[wb.sheetnames[-1]]

            target_ws = wb.copy_worksheet(source_ws)
            final_title = nombre_pestana
            c = 1
            while final_title in wb.sheetnames:
                final_title = f"{nombre_pestana}_{c}"
                c += 1
            target_ws.title = final_title
            nombre_pestana = final_title

        # 1. Cabecera
        target_ws['C5'] = cliente
        target_ws['C6'] = float(ruc) if ruc.isdigit() else ruc
        target_ws['C7'] = planta
        target_ws['C8'] = responsable
        target_ws['C9'] = liquidacion
        target_ws['C10'] = fecha_emision

        # 2. Limpiar filas de servicios previas (filas 13 a 35)
        for r in range(13, 36):
            for col in range(1, 9):
                target_ws.cell(r, col).value = None

        # 3. Escribir items con fórmulas
        fila_inicio = 13
        for i, it in enumerate(items):
            r = fila_inicio + i
            num_v = str(it.get("num", "")).strip()
            desc_v = str(it.get("descripcion", "")).strip()
            fec_v = str(it.get("fecha", "-")).strip()
            guia_v = str(it.get("guia", "-")).strip()
            cant_v = limpiar_monto_numerico(it.get("cantidad", 1))
            costo_v = limpiar_monto_numerico(it.get("costo", 0))

            target_ws.cell(r, 1).value = float(num_v) if num_v.isdigit() else (num_v if num_v else None)
            target_ws.cell(r, 2).value = desc_v
            target_ws.cell(r, 4).value = fec_v
            target_ws.cell(r, 5).value = guia_v
            target_ws.cell(r, 6).value = cant_v
            target_ws.cell(r, 7).value = costo_v
            target_ws.cell(r, 8).value = f"=F{r}*G{r}"

        # 4. Bloque de Totales
        n_items = len(items) if items else 1
        ultima_fila_items = fila_inicio + n_items - 1
        fila_total_mes = ultima_fila_items + 1
        fila_subtotal = fila_total_mes + 2
        fila_igv = fila_subtotal + 1
        fila_total = fila_igv + 1
        fila_detraccion = fila_total + 1

        target_ws.cell(fila_total_mes, 1).value = "TOTAL GENERADO EN EL MESES"
        target_ws.cell(fila_total_mes, 8).value = f"=SUM(H{fila_inicio}:H{ultima_fila_items})"

        target_ws.cell(fila_subtotal, 1).value = "* Metodo de pago: transferencia."
        target_ws.cell(fila_subtotal, 7).value = "SUB TOTAL"
        target_ws.cell(fila_subtotal, 8).value = f"=H{fila_total_mes}"

        target_ws.cell(fila_igv, 1).value = "* Forma de pago: Al contado."
        target_ws.cell(fila_igv, 7).value = "IGV"
        target_ws.cell(fila_igv, 8).value = f"=H{fila_total}-H{fila_subtotal}"

        target_ws.cell(fila_total, 1).value = "* Numero de cuenta Corriente soles en Interbank es: 200-3005602805. "
        target_ws.cell(fila_total, 7).value = "TOTAL"
        target_ws.cell(fila_total, 8).value = f"=H{fila_subtotal}*1.18"

        target_ws.cell(fila_detraccion, 1).value = "* Numero de cuenta de detracción: 00050086984"

        # Firmas al pie
        fila_firmas = fila_detraccion + 9
        target_ws.cell(fila_firmas, 3).value = "______________________________________"
        target_ws.cell(fila_firmas, 5).value = "_______________________________________"
        target_ws.cell(fila_firmas + 1, 3).value = "EPMI S.A.C."
        target_ws.cell(fila_firmas + 1, 5).value = cliente

        # 5. Guardar en Google Drive
        subir_workbook_drive(drive, wb)

        return {
            "success": True,
            "message": f"Pestaña '{nombre_pestana}' guardada con éxito en Google Sheets.",
            "sheet_name": nombre_pestana,
            "gid": nombre_pestana,
            "cliente": cliente,
            "liquidacion": liquidacion
        }
    except Exception as e:
        logger.error(f"Error guardando valorización: {e}", exc_info=True)
        return {
            "success": False,
            "error": str(e)
        }

def exportar_valorizacion_pdf(sheet_name, nombre_archivo=None):
    """
    Exporta la pestaña seleccionada del archivo Excel a PDF nativo,
    guardándolo en la carpeta de Drive designada y retornando bytes y link.
    """
    try:
        creds = obtener_credenciales()
        if (creds.expired or not creds.valid) and creds.refresh_token:
            creds.refresh(Request())

        drive = build('drive', 'v3', credentials=creds)

        # 1. Descargar el archivo Excel actual
        request = drive.files().get_media(fileId=SHEET_VALORIZACIONES_ID)
        fh = io.BytesIO()
        downloader = MediaIoBaseDownload(fh, request)
        done = False
        while not done:
            _, done = downloader.next_chunk()
        fh.seek(0)

        # 2. Subir como Google Sheet temporal para conversión perfecta a PDF
        media = MediaIoBaseUpload(
            fh,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            resumable=True
        )
        temp_body = {
            'name': f'TMP_CONVERT_{int(time.time())}',
            'mimeType': 'application/vnd.google-apps.spreadsheet',
            'parents': [DRIVE_FOLDER_VALORIZACIONES_PDF]
        }
        temp_file = drive.files().create(body=temp_body, media_body=media, fields='id').execute()
        temp_id = temp_file['id']

        # 3. Obtener el GID de la hoja solicitada
        import gspread
        import requests
        gc = gspread.authorize(creds)
        temp_doc = gc.open_by_key(temp_id)

        try:
            ws = temp_doc.worksheet(str(sheet_name))
        except gspread.exceptions.WorksheetNotFound:
            ws = temp_doc.worksheets()[-1]

        target_gid = ws.id

        # 4. Descargar PDF renderizado
        url_export = (
            f"https://docs.google.com/spreadsheets/d/{temp_id}/export"
            f"?format=pdf&gid={target_gid}&size=A4&portrait=true&fitw=true&gridlines=false"
        )
        headers = {'Authorization': f'Bearer {creds.token}'}
        resp = requests.get(url_export, headers=headers, timeout=40)

        # 5. Eliminar archivo temporal
        try:
            drive.files().delete(fileId=temp_id).execute()
        except Exception:
            pass

        if resp.status_code != 200:
            raise ValueError(f"Error descargando PDF: status {resp.status_code}")

        pdf_bytes = resp.content

        if not nombre_archivo:
            nombre_archivo = f"Valorizacion_{sheet_name}.pdf"
        if not nombre_archivo.lower().endswith(".pdf"):
            nombre_archivo += ".pdf"

        # 6. Subir PDF final a la carpeta de PDFs
        meta_pdf = {
            'name': nombre_archivo,
            'parents': [DRIVE_FOLDER_VALORIZACIONES_PDF],
            'mimeType': 'application/pdf'
        }
        media_pdf = MediaInMemoryUpload(pdf_bytes, mimetype='application/pdf', resumable=True)
        final_pdf = drive.files().create(
            body=meta_pdf,
            media_body=media_pdf,
            fields='id, name, webViewLink'
        ).execute()

        drive_link = final_pdf.get('webViewLink', '')
        file_id = final_pdf.get('id', '')
        logger.info(f"✅ PDF de Valorización guardado en Drive: {nombre_archivo} ({drive_link})")

        return {
            "success": True,
            "pdf_bytes": pdf_bytes,
            "drive_link": drive_link,
            "file_id": file_id,
            "filename": nombre_archivo
        }
    except Exception as e:
        logger.error(f"Error exportando valorización a PDF: {e}", exc_info=True)
        return {
            "success": False,
            "error": str(e)
        }
